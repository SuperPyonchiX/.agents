"""監査で見つけた誤合格・途中失敗の回帰検査。外部通信は全てモック。

リポジトリ直下で python tools/test_skill_audit_regressions.py を実行する。
標準ライブラリのみ。終了コード0=全件成功、1=テスト失敗。
"""
import argparse
import contextlib
import importlib.util
import http.client
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]


def load(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'skills' / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


trace = load('trace', 'nextdesign-cpp14-implementation/scripts/check_traceability.py')
audit = load('audit', 'skill-portfolio-audit/scripts/audit_skills.py')
claude = load('ask_claude', 'second-opinion/scripts/ask_claude.py')
agy = load('ask_antigravity', 'second-opinion/scripts/ask_antigravity.py')
codex = load('ask_codex', 'second-opinion/scripts/ask_codex.py')
sys.path.insert(0, str(ROOT / 'skills/notion-api/scripts'))
import notion_page as notion
recalc = load('recalc_check', 'operate-excel/scripts/recalc_check.py')
EXCEL = ROOT/'skills/operate-excel'
HAS_OPENPYXL = importlib.util.find_spec('openpyxl') is not None
PPT = ROOT/'skills/operate-powerpoint'
HAS_PPTX = importlib.util.find_spec('pptx') is not None
WORD = ROOT/'skills/operate-word'
HAS_DOCX = importlib.util.find_spec('docx') is not None
PDF = ROOT/'skills/operate-pdf'
HAS_PYPDF = importlib.util.find_spec('pypdf') is not None
HAS_PDFIUM = importlib.util.find_spec('pypdfium2') is not None and importlib.util.find_spec('PIL') is not None
DRAWIO = ROOT/'skills/drawio-diagram'


class TemporaryCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def file(self, name, text):
        p = self.folder / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding='utf-8')
        return p

    def run_ledger(self, skill, script, text, expected, diagnostic=None):
        p = self.file('ledger.md', text)
        result = subprocess.run([sys.executable, str(ROOT/'skills'/skill/'scripts'/script), str(p)],
                                capture_output=True, text=True, encoding='utf-8',
                                env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        self.assertNotIn('Traceback', result.stderr)
        if expected == 1:
            self.assertIn('ERROR', result.stdout)
        if diagnostic:
            self.assertIn(diagnostic, result.stdout)
        return result.stdout


class LedgerTests(TemporaryCase):
    def test_review_missing_reason_column(self):
        self.run_ledger('cpp14-code-review', 'check_review_log.py',
                        '| ID | 状態 |\n|---|---|\n| R-01 | accepted |\n', 1, '必須の対応/理由列がない')

    def test_review_valid_and_duplicate(self):
        text = '| ID | 状態 | 対応/理由 |\n|---|---|---|\n| R-01 | accepted | 了承済み |\n'
        self.run_ledger('cpp14-code-review', 'check_review_log.py', text, 0)
        self.run_ledger('cpp14-code-review', 'check_review_log.py', text + '| R-01 | fixed | 修正済み |\n', 1, '重複')

    def test_review_short_row_after_valid(self):
        text = '| ID | 状態 | 対応/理由 |\n|---|---|---|\n| R-01 | fixed | 修正済 |\n| R-02 | open |\n'
        self.run_ledger('cpp14-code-review', 'check_review_log.py', text, 1, '列数が見出しと一致しない')

    def test_defect_missing_columns(self):
        self.run_ledger('cpp14-defect-analysis', 'check_defect_log.py',
                        '| ID | 状態 |\n|---|---|\n| H-01 | confirmed |\n', 1, '必須列がない')

    def test_defect_valid_and_missing_result(self):
        text = '| ID | 状態 | 検証方法 | 検証結果 |\n|---|---|---|---|\n| H-01 | confirmed | ログ観測 | 再現 |\n'
        self.run_ledger('cpp14-defect-analysis', 'check_defect_log.py', text, 0)
        self.run_ledger('cpp14-defect-analysis', 'check_defect_log.py', text.replace('再現', ''), 1, '検証結果が空')

    def test_trace_statuses(self):
        for status, valid in [('fixed', True), ('accepted', True), ('opne', False),
                              ('逸脱承認待ち', False), ('pending', False), ('open', False)]:
            with self.subTest(status=status):
                p = self.file('trace.md', '| ID | 状態 |\n|---|---|\n| R-01 | ' + status + ' |\n')
                rep = trace.Report()
                counts = trace.check_review_log(p, rep)
                self.assertEqual(not rep.errors, valid, rep.errors)
                self.assertEqual(counts['closed'], int(valid))

    def test_trace_empty_or_duplicate(self):
        for rows in ['', '| R-01 | fixed |\n| R-01 | fixed |\n']:
            p = self.file('trace.md', '| ID | 状態 |\n|---|---|\n' + rows)
            rep = trace.Report()
            trace.check_review_log(p, rep)
            self.assertTrue(rep.errors)

    def test_ui_short_row_after_valid(self):
        valid = '| UI-001 | 指摘 | 対象 | 前 | 変更 | 後 | before.png | after.png | 1 | deferred | 利用者が保留 |\n'
        self.run_ledger('ui-visual-verify', 'check_ui_ledger.py', valid, 0)
        self.run_ledger('ui-visual-verify', 'check_ui_ledger.py', valid + '| UI-002 | open |\n', 1, '列数が不正')
        self.run_ledger('ui-visual-verify', 'check_ui_ledger.py', valid + valid, 1, 'IDが重複')

    def test_ui_closed_requires_evidence(self):
        text = '| UI-001 | 指摘 | 対象 | 前 | 変更 | 120px | before.png | after.png | 1 | closed | 確認済 |\n'
        self.run_ledger('ui-visual-verify', 'check_ui_ledger.py', text, 1)
        self.file('after.png', 'fixture')
        self.run_ledger('ui-visual-verify', 'check_ui_ledger.py', text, 0)


class AuditTests(TemporaryCase):
    def test_warning_count_summary_and_full_output(self):
        validator = self.file('validator.py', "print('WARN   first: detail')\n" +
                              "print('\\n'.join('OK detail' for _ in range(10)))\nprint('ERROR 0 件 / WARN 1 件')\n")
        result = audit.run_validator(str(validator), str(self.folder), ['demo'])['demo']
        self.assertEqual((result['errors'], result['warns'], result['exit_code']), (0, 1, 0))
        self.assertTrue(result['output'][0].startswith('WARN'))
        self.assertGreater(len(result['output']), 6)

    def test_warning_is_aggregated_even_exit_zero(self):
        skills = self.folder/'skills'
        self.file('skills/demo/SKILL.md', '---\nname: demo\ndescription: 調査するときに使う。\n---\n')
        validator = self.file('validator.py', "print('WARN   example: missing reference')\nprint('ERROR 0 件 / WARN 1 件')\n")
        output = self.folder/'out.json'
        with contextlib.redirect_stdout(io.StringIO()):
            code = audit.main([str(skills), '--validator', str(validator), '-o', str(output)])
        result = json.loads(output.read_text(encoding='utf-8'))
        self.assertEqual(code, 0)
        self.assertEqual(len(result['warns']), 1)
        self.assertEqual(result['validation']['demo']['errors'], 0)

    def test_content_hash_and_deletion_resume(self):
        skills = self.folder/'skills'
        self.file('skills/demo/SKILL.md', '---\nname: demo\ndescription: 調査するときに使う。\n---\n')
        doc = self.file('skills/demo/references/info.md', 'before')
        removed = self.file('skills/demo/references/removed.md', 'delete me')
        self.file('skills/demo/local/secret.txt', 'do not hash')
        old = audit.fingerprint(str(skills), ['demo'])
        self.assertFalse(any('/local/' in x for x in old))
        stamp = doc.stat().st_mtime_ns
        doc.write_text('after!', encoding='utf-8')
        os.utime(doc, ns=(stamp, stamp))
        removed.unlink()
        previous = self.file('old.json', json.dumps({'fingerprints': old}))
        output = self.folder/'new.json'
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(audit.main([str(skills), '--compare', str(previous), '-o', str(output)]), 0)
        changed = json.loads(output.read_text(encoding='utf-8'))['changed_files']
        self.assertEqual(changed, ['demo/references/info.md', 'demo/references/removed.md'])


class NotionTests(TemporaryCase):
    def args(self, count=201):
        return argparse.Namespace(properties='{}', blocks=json.dumps([{'type': 'paragraph'}]*count),
                                  data_source_id='fake', icon=None, progress=str(self.folder/'progress.json'))

    def invoke(self, method, args, fake):
        stdout, stderr = io.StringIO(), io.StringIO()
        code = 0
        with patch.object(notion, 'request', side_effect=fake) as req, \
                contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            try:
                method(args)
            except SystemExit as exc:
                code = exc.code
        raw = stdout.getvalue()
        return code, json.loads(raw) if raw else None, req.call_args_list

    def test_create_three_batches_success(self):
        sizes = []
        def fake(method, path, payload):
            sizes.append(len(payload['children']))
            return {'id': 'page', 'url': 'https://example.invalid/page'}
        code, result, calls = self.invoke(notion.cmd_create, self.args(), fake)
        self.assertEqual((code, result['status'], result['confirmed_blocks']), (0, 'complete', 202))
        self.assertEqual(sizes, [100, 100, 2])
        self.assertEqual(len(calls), 3)

    def test_partial_create_preserves_id_and_confirmed_progress(self):
        counter = 0
        def fake(method, path, payload):
            nonlocal counter
            counter += 1
            if counter == 3:
                raise SystemExit(1)
            return {'id': 'page', 'url': 'https://example.invalid/page'}
        code, result, calls = self.invoke(notion.cmd_create, self.args(), fake)
        self.assertEqual(code, 3)
        self.assertEqual((result['id'], result['confirmed_blocks'], result['pending_blocks']), ('page', 200, 2))
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(json.loads((self.folder/'progress.json').read_text(encoding='utf-8')), result)
        self.assertEqual(len(calls), 3)  # 不明のバッチを自動再送しない

    def test_create_first_patch_failure_preserves_created_id(self):
        def fake(method, path, payload):
            if method == 'PATCH':
                saved = json.loads((self.folder/'progress.json').read_text(encoding='utf-8'))
                self.assertEqual(saved['id'], 'page')
                raise OSError('timeout')
            return {'id': 'page'}
        code, result, _ = self.invoke(notion.cmd_create, self.args(101), fake)
        self.assertEqual((code, result['confirmed_blocks'], result['id']), (3, 100, 'page'))

    def test_post_response_unknown_does_not_retry(self):
        code, result, calls = self.invoke(notion.cmd_create, self.args(0), OSError('timeout'))
        self.assertEqual((code, result['status'], result['id']), (3, 'unknown', None))
        self.assertEqual(len(calls), 1)

    def test_missing_token_keeps_input_error(self):
        code, result, _ = self.invoke(notion.cmd_create, self.args(0), SystemExit(2))
        self.assertEqual((code, result), (2, None))

    def test_unreadable_post_response_is_unknown(self):
        for fake in [ValueError('invalid response JSON'), http.client.IncompleteRead(b'partial'), lambda *a: {}]:
            code, result, calls = self.invoke(notion.cmd_create, self.args(0), fake)
            self.assertEqual((code, result['status']), (3, 'unknown'))
            self.assertEqual(len(calls), 1)

    def test_http_exception_during_patch_preserves_id(self):
        def fake(method, path, payload):
            if method == 'PATCH':
                raise http.client.BadStatusLine('broken HTTP response')
            return {'id': 'page'}
        code, result, calls = self.invoke(notion.cmd_create, self.args(101), fake)
        self.assertEqual((code, result['status'], result['id']), (3, 'partial', 'page'))
        self.assertEqual(len(calls), 2)

    def test_interrupt_leaves_recoverable_progress(self):
        def fake(method, path, payload):
            if method == 'PATCH':
                raise KeyboardInterrupt()
            return {'id': 'page'}
        with self.assertRaises(KeyboardInterrupt):
            self.invoke(notion.cmd_create, self.args(101), fake)
        saved = json.loads((self.folder/'progress.json').read_text(encoding='utf-8'))
        self.assertEqual((saved['status'], saved['id'], saved['confirmed_blocks'], saved['pending_blocks']),
                         ('sending', 'page', 100, 2))

    def test_progress_unwritable_stops_before_post(self):
        args = self.args()
        args.progress = str(self.folder/'missing'/'state.json')
        code, result, calls = self.invoke(notion.cmd_create, args, AssertionError('no request expected'))
        self.assertEqual((code, result, len(calls)), (2, None, 0))

    def test_progress_write_fails_after_post_id_still_returned(self):
        def fake(method, path, payload):
            return {'id': 'page'}
        real_save = notion.save_progress
        def save(args, state):
            if state['id']:
                raise OSError('disk full')
            real_save(args, state)
        with patch.object(notion, 'save_progress', side_effect=save):
            code, result, calls = self.invoke(notion.cmd_create, self.args(0), fake)
        self.assertEqual((code, result['id'], result['confirmed_blocks']), (3, 'page', 1))
        self.assertEqual(len(calls), 1)

    def test_append_mismatch_prevents_write(self):
        args = self.args(2)
        args.page_id, args.expected_count = 'page', 99
        code, result, calls = self.invoke(notion.cmd_append, args, lambda *a: {'results': [1]*100, 'has_more': False})
        self.assertEqual((code, result, len(calls)), (2, None, 1))
        self.assertEqual(calls[0].args[0], 'GET')

    def test_append_reads_all_pages_then_sends_only_missing(self):
        args = self.args(2)
        args.page_id, args.expected_count = 'page', 101
        responses = iter([{'results': [1]*100, 'has_more': True, 'next_cursor': 'next'},
                          {'results': [1], 'has_more': False}, {}])
        code, result, calls = self.invoke(notion.cmd_append, args, lambda *a: next(responses))
        self.assertEqual((code, result['confirmed_blocks']), (0, 103))
        self.assertEqual([call.args[0] for call in calls], ['GET', 'GET', 'PATCH'])
        self.assertEqual(len(calls[-1].args[2]['children']), 2)


class ClaudeTests(TemporaryCase):
    def test_normal_tools_and_no_delegation(self):
        cmd = claude.build_command('claude')
        for option in ['--tools', '--safe-mode', '--strict-mcp-config', '--disable-slash-commands', '--permission-mode', '--system-prompt']:
            self.assertNotIn(option, cmd)
        self.assertEqual(cmd[cmd.index('--disallowedTools')+1], 'Agent,Task')
        self.assertIn('別の生成AIへの相談', cmd[cmd.index('--append-system-prompt')+1])

    def test_success_and_cli_error(self):
        q = self.file('q.md', '根拠を示して検討してください。')
        out = self.folder/'answer.md'
        for success in [True, False]:
            with self.subTest(success=success):
                response = subprocess.CompletedProcess([], 0, json.dumps({'result': '独立した回答', 'is_error': not success}), '')
                with patch.object(claude.shutil, 'which', return_value='claude'), \
                        patch.object(claude.subprocess, 'run', return_value=response) as run, \
                        contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    code = claude.main(['--prompt-file', str(q), '--out', str(out), '--cd', str(self.folder)])
                self.assertEqual(code, 0 if success else 1)
                self.assertEqual(Path(run.call_args.kwargs['cwd']), self.folder)
                if success:
                    self.assertEqual(out.read_text(encoding='utf-8').strip(), '独立した回答')


class AntigravityTests(TemporaryCase):
    def test_command_is_headless_without_skip_permissions(self):
        cmd = agy.build_command('agy', 600)
        self.assertEqual(cmd[1], '-p')
        self.assertIn('--disable-slash-commands', cmd)
        self.assertEqual(cmd[cmd.index('--output-format')+1], 'json')
        self.assertNotIn('--dangerously-skip-permissions', cmd)
        self.assertIn('別の生成AIへの相談', cmd[2])

    def test_runs_on_copy_and_keeps_original(self):
        src = self.folder/'src'
        self.file('src/a.txt', '元の内容')
        self.file('src/.git/HEAD', 'ref')
        q = self.file('q.md', '根拠を示して検討してください。')
        out = self.folder/'answer.md'
        for status in ['SUCCESS', 'ERROR']:
            with self.subTest(status=status):
                def fake_run(cmd, **kwargs):
                    ws = Path(kwargs['cwd'])
                    self.assertNotEqual(ws, src)
                    self.assertFalse((ws/'.git').exists())
                    self.assertIn('根拠を示して', (ws/agy.QUESTION_NAME).read_text(encoding='utf-8'))
                    (ws/'a.txt').write_text('書き換え', encoding='utf-8')
                    return subprocess.CompletedProcess(cmd, 0, json.dumps({'status': status, 'response': '独立した回答'}), '')
                with patch.object(agy.shutil, 'which', return_value='agy'), \
                        patch.object(agy.subprocess, 'run', side_effect=fake_run), \
                        contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    code = agy.main(['--prompt-file', str(q), '--out', str(out), '--cd', str(src)])
                self.assertEqual(code, 0 if status == 'SUCCESS' else 1)
                self.assertEqual((src/'a.txt').read_text(encoding='utf-8'), '元の内容')
                self.assertEqual(out.exists(), status == 'SUCCESS')

    def test_missing_cli(self):
        q = self.file('q.md', '質問')
        with patch.object(agy.shutil, 'which', return_value=None), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(agy.main(['--prompt-file', str(q), '--out', str(self.folder/'a.md')]), 2)


class ConsoleEncodingTests(unittest.TestCase):
    """PowerShell はパイプの出力をコンソール出力コードページで読むので、それに合わせて書く。"""
    TEXT = '回答保存: C:\\テスト\\a0.md'

    def write(self, module, cp, env=None, start='utf-8'):
        streams = [io.TextIOWrapper(io.BytesIO(), encoding=start) for _ in range(2)]
        windll = type('W', (), {'kernel32': type('K', (), {'GetConsoleOutputCP': staticmethod(lambda: cp)})})
        with patch.object(module.sys, 'platform', 'win32'), patch.dict(os.environ, env or {}, clear=True), \
                patch('ctypes.windll', windll, create=True), \
                patch.object(module.sys, 'stdout', streams[0]), patch.object(module.sys, 'stderr', streams[1]):
            module.match_console_encoding()
        for stream in streams:
            stream.write(self.TEXT + '\U0001F600')
            stream.flush()
        return [(stream.encoding, stream.buffer.getvalue()) for stream in streams]

    def test_follows_console_code_page(self):
        for module in [agy, claude, codex]:
            for cp, expected in [(932, 'cp932'), (65001, 'utf-8')]:
                with self.subTest(module=module.__name__, cp=cp):
                    for encoding, data in self.write(module, cp):
                        self.assertEqual(encoding, expected)
                        # cp932 で表せない文字は落とさず \\U... で残す
                        self.assertTrue(data.decode(expected).startswith(self.TEXT))

    def test_respects_pythonioencoding_and_missing_console(self):
        for module in [agy, claude, codex]:
            with self.subTest(module=module.__name__):
                self.assertEqual(self.write(module, 932, {'PYTHONIOENCODING': 'utf-8'})[0][0], 'utf-8')
                self.assertEqual(self.write(module, 0, start='cp932')[0][0], 'cp932')


class ExcelTests(TemporaryCase):
    WORKBOOK = ('<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<sheets><sheet name="表" sheetId="1" r:id="rId1"/></sheets></workbook>')
    RELS = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="worksheet" Target="worksheets/sheet1.xml"/></Relationships>')

    def book(self, cells):
        p = self.folder/'calc.xlsx'
        with zipfile.ZipFile(p, 'w') as z:
            z.writestr('xl/workbook.xml', self.WORKBOOK)
            z.writestr('xl/_rels/workbook.xml.rels', self.RELS)
            z.writestr('xl/worksheets/sheet1.xml',
                       '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                       '<sheetData><row r="1">%s</row></sheetData></worksheet>' % cells)
        return p

    def run_script(self, script, *args):
        return subprocess.run([sys.executable, str(EXCEL/'scripts'/script), *map(str, args)],
                              capture_output=True, text=True, encoding='utf-8',
                              env=dict(os.environ, PYTHONIOENCODING='utf-8'))

    def test_recalc_scan_counts_blank_result_and_array_and_shared(self):
        p = self.book('<c r="A1" t="str"><f>IF(B1&gt;0,B1,"")</f><v></v></c>'
                      '<c r="B1"><f t="array" ref="B1">SUM(C1:C2*2)</f><v>4</v></c>'
                      '<c r="C1" t="e"><f t="shared" ref="C1:D1" si="0">1/0</f><v>#DIV/0!</v></c>'
                      '<c r="D1" t="e"><f t="shared" si="0"/><v>#DIV/0!</v></c>'
                      '<c r="E1"><f>F1*2</f></c>'
                      '<c r="F1" t="inlineStr"><is><t>#N/A</t></is></c>')
        formulas, uncalculated, errors, summary = recalc.scan(p)
        self.assertEqual(formulas, 5)
        self.assertEqual(uncalculated, 1)  # E1 だけ。空文字を返す A1 は計算済み
        self.assertEqual(summary, {'#DIV/0!': 2})  # 文字列の "#N/A" はエラーではない
        self.assertEqual([e['formula'] for e in errors], ['=1/0', '（C1 の共有数式）'])

    def test_strip_cached_values_keeps_constants(self):
        src = self.book('<c r="A1"><v>5</v></c><c r="B1"><f>A1*2</f><v>999</v></c>'
                        '<c r="C1"><f t="shared" si="0"/><v>1</v></c>')
        dst = self.folder/'stripped.xlsx'
        recalc.strip_cached_values(src, dst)
        with zipfile.ZipFile(dst) as z:
            xml = z.read('xl/worksheets/sheet1.xml').decode()
        self.assertIn('<c r="A1"><v>5</v></c>', xml)
        self.assertIn('<c r="B1"><f>A1*2</f></c>', xml)
        self.assertIn('<c r="C1"><f t="shared" si="0"/></c>', xml)

    @unittest.skipUnless(recalc.find_soffice(), 'LibreOffice が無い')
    def test_recalc_ignores_stale_cached_value(self):
        src = EXCEL/'evals/files/stale_cache.xlsx'
        before = src.read_bytes()
        result = self.run_script('recalc_check.py', src)
        self.assertEqual(result.returncode, 1, result.stderr)
        out = json.loads(result.stdout)
        self.assertEqual([(e['cell'], e['error']) for e in out['errors']], [('D4', '#DIV/0!')])
        self.assertEqual(src.read_bytes(), before)

    @unittest.skipUnless(HAS_OPENPYXL, 'openpyxl が無い')
    def test_inspect_unknown_sheet_is_usage_error(self):
        result = self.run_script('inspect_xlsx.py', EXCEL/'evals/files/stale_cache.xlsx', '--sheet', '無い')
        self.assertEqual(result.returncode, 2)
        self.assertNotIn('Traceback', result.stderr)

    @unittest.skipUnless(HAS_OPENPYXL, 'openpyxl が無い')
    def test_inspect_and_diff_on_bloated_sheet(self):
        import openpyxl
        src = EXCEL/'evals/files/bloated_blank_if.xlsx'
        result = self.run_script('inspect_xlsx.py', src)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('A1:C4（4 行 × 3 列）', result.stdout)
        self.assertIn('使用範囲は A1:GR100000', result.stdout)
        after = self.folder/'after.xlsx'
        wb = openpyxl.load_workbook(src)
        wb.active['B3'] = 65
        wb.save(after)
        result = self.run_script('diff_xlsx.py', src, after)
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertIn('変わったセル: 1 件（範囲 B3:B3）', result.stdout)
        self.assertIn('| B3 | 値 | （空） | 65 |', result.stdout)


@unittest.skipUnless(HAS_PPTX, 'python-pptx が無い')
class PowerPointTests(TemporaryCase):
    DECK = PPT/'evals/files/q3_report.pptx'

    def run_script(self, *args):
        return subprocess.run([sys.executable, str(PPT/'scripts/inspect_pptx.py'), *map(str, args)],
                              capture_output=True, text=True, encoding='utf-8',
                              env=dict(os.environ, PYTHONIOENCODING='utf-8'))

    def helpers(self):
        import re
        md = (PPT/'references/python-pptx-helpers.md').read_text(encoding='utf-8')
        blocks = re.findall(r"```python\n(.*?)```", md, re.S)
        g = {}
        exec(blocks[0], g)
        return g

    def test_check_finds_empty_placeholder_table_overlap_and_hidden_slide(self):
        result = self.run_script(self.DECK, '--check')
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertIn('スライド 1（画像 p1）: 「Subtitle 2」が空のプレースホルダ', result.stdout)
        self.assertIn('スライド 4（画像 p3）: 「Table 2」と「TextBox 3」が重なっている', result.stdout)
        self.assertIn('スライド 5（画像 p4）: 「Content Placeholder 2」の文字があふれている疑い', result.stdout)
        self.assertIn('非表示スライド: 3', result.stdout)
        self.assertNotIn('スライド 2', result.stdout)  # 収まっているスライドは何も出さない
        self.assertNotIn('ph[', result.stdout)  # --check だけなら図形一覧は出さない

    def test_full_shows_table_cells(self):
        result = self.run_script(self.DECK, '--slides', '4', '--full')
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('3 | ログが重複する | 11/19', result.stdout)

    def test_bad_slide_range_is_usage_error(self):
        for spec in ('a', '9', '3-1', ',', '1-', '1,,2'):
            result = self.run_script(self.DECK, '--slides', spec)
            self.assertEqual(result.returncode, 2, spec)
            self.assertNotIn('Traceback', result.stderr)

    def test_delete_then_add_keeps_every_slide(self):
        g = self.helpers()
        prs = g['Presentation'](str(self.DECK))
        g['delete_slide'](prs, 2)
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = '次回の報告は12月'
        out = self.folder/'out.pptx'
        prs.save(str(out))
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
        self.assertEqual(len(names), len(set(names)))
        titles = [s.shapes.title.text for s in g['Presentation'](str(out)).slides]
        self.assertEqual(titles, ['第3四半期の不具合報告', '不具合は 3 件に減った', '残課題の一覧',
                                  '再発防止の取り組み', '次回の報告は12月'])

    def check(self, prs):
        out = self.folder/'case.pptx'
        prs.save(str(out))
        result = self.run_script(out, '--check')
        self.assertNotIn('Traceback', result.stderr)
        return result

    def test_check_edge_cases(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Mm, Pt
        prs = Presentation()
        blank = prs.slide_layouts[6]
        # 1: 段落に指定した 10pt（run には指定なし）と、空段落だけが 10pt の枠
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Mm(10), Mm(10), Mm(100), Mm(20))
        tb.text_frame.text = '段落で小さくした'
        tb.text_frame.paragraphs[0].font.size = Pt(10)
        tb2 = s.shapes.add_textbox(Mm(10), Mm(60), Mm(100), Mm(30))
        tb2.text_frame.text = '本文'
        tb2.text_frame.add_paragraph().font.size = Pt(10)
        # 2: 表のセルの 8pt
        s = prs.slides.add_slide(blank)
        cell = s.shapes.add_table(1, 1, Mm(10), Mm(10), Mm(60), Mm(20)).table.cell(0, 0)
        cell.text = '細かい表'
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        # 3: 折り返さない枠に4段落（縦にあふれる）
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Mm(10), Mm(10), Mm(100), Mm(10))
        tb.text_frame.text = '一'
        for _ in range(3):
            tb.text_frame.add_paragraph().text = '二'
        # 4: グラフと画像の重なり
        s = prs.slides.add_slide(blank)
        data = CategoryChartData()
        data.categories = ['a', 'b']
        data.add_series('x', (1, 2))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Mm(10), Mm(10), Mm(80), Mm(60), data).name = 'グラフ'
        png = self.folder/'dot.png'
        import base64
        png.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFBQIAX8jx0gAAAABJRU5ErkJggg=='))
        s.shapes.add_picture(str(png), Mm(50), Mm(30), Mm(60), Mm(40)).name = '写真'
        # 5: 同じグループの中の重なりは見ない／90度回転したグループの子はスライド外へ出る
        s = prs.slides.add_slide(blank)
        g = s.shapes.add_group_shape()
        card = g.shapes.add_textbox(Mm(150), Mm(20), Mm(80), Mm(40))
        card.text_frame.text = 'カードの見出し'
        inner = g.shapes.add_textbox(Mm(160), Mm(30), Mm(60), Mm(20))
        inner.text_frame.text = 'カードの本文'
        g2 = s.shapes.add_group_shape()
        long = g2.shapes.add_textbox(Mm(10), Mm(170), Mm(200), Mm(10))
        long.text_frame.text = '横長の帯'
        long.name = '帯'
        g2._element.grpSpPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm').set('rot', str(90 * 60000))
        out = self.check(prs).stdout
        self.assertIn('スライド 1: 「TextBox 1」に 10pt の文字がある', out)
        self.assertNotIn('「TextBox 2」に', out)  # 空段落の 10pt は数えない
        self.assertIn('スライド 2: 「Table 1」に 8pt の文字がある', out)
        self.assertIn('スライド 3: 「TextBox 1」の文字があふれている疑い（推定: 必要な高さ', out)
        self.assertIn('「グラフ」と「写真」が重なっている', out)
        self.assertNotIn('カード', out)
        self.assertIn('スライド 5: 「帯」がスライドの外にはみ出している', out)

    def test_delete_linked_slide_removes_link(self):
        g = self.helpers()
        prs = g['Presentation']()
        slides = [prs.slides.add_slide(prs.slide_layouts[5]) for _ in range(3)]
        for i, s in enumerate(slides):
            s.shapes.title.text = 's%d' % i
        slides[0].shapes.title.click_action.target_slide = slides[1]
        g['delete_slide'](prs, 1)
        out = self.folder/'link.pptx'
        prs.save(str(out))
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            first = z.read('ppt/slides/slide1.xml').decode()
        self.assertEqual(len(names), len(set(names)))
        self.assertNotIn('hlinkClick', first)
        self.assertEqual([s.shapes.title.text for s in g['Presentation'](str(out)).slides], ['s0', 's2'])

    def test_potx_opens(self):
        potx = self.folder/'t.potx'
        with zipfile.ZipFile(self.DECK) as src, zipfile.ZipFile(potx, 'w') as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == '[Content_Types].xml':
                    data = data.replace(b'presentationml.presentation.main+xml', b'presentationml.template.main+xml')
                dst.writestr(item, data)
        self.assertEqual(len(self.helpers()['open_template'](str(potx)).slides), 5)
        result = self.run_script(potx, '--outline')
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('スライド 4（Title Only・画像 p3）: 残課題の一覧', result.stdout)


@unittest.skipUnless(HAS_DOCX, 'python-docx が無い')
class WordTests(TemporaryCase):
    MANUAL = WORD/'evals/files/manual.docx'

    def run_script(self, name, *args):
        return subprocess.run([sys.executable, str(WORD/'scripts'/name), *map(str, args)],
                              capture_output=True, text=True, encoding='utf-8',
                              env=dict(os.environ, PYTHONIOENCODING='utf-8'))

    def replace(self, src, *args):
        out = self.folder/'out.docx'
        if out.exists():
            out.unlink()
        return self.run_script('docx_replace.py', src, out, *args), out

    def test_inspect_reads_hidden_text(self):
        result = self.run_script('inspect_docx.py', self.MANUAL)
        self.assertEqual(result.returncode, 0, result.stderr)
        for text in ('対象機種: Alpha 2',                     # コンテンツコントロールの中
                     '[テキストボックス] Alpha 2 クイックガイド',
                     '（Alpha 2 の後継機を除く）',            # 変更履歴で挿入した文字
                     '[^2]: Alpha 2 は2024年4月に発売した。',
                     '先頭ページのヘッダー: Alpha 2 取扱説明書',
                     '| 項目 | ← | 値 |',                    # 横に結合したセルを繰り返さない
                     'テキストボックス 1、脚注 1'):
            self.assertIn(text, result.stdout)
        outline = self.run_script('inspect_docx.py', self.MANUAL, '--outline').stdout
        self.assertIn('§3 ## 対応機種', outline)
        part = self.run_script('inspect_docx.py', self.MANUAL, '--section', '2').stdout
        self.assertIn('AC 100 V', part)
        self.assertNotIn('クイックガイド', part)
        bad = self.run_script('inspect_docx.py', self.MANUAL, '--section', '9')
        self.assertEqual(bad.returncode, 2)

    def test_long_body_prints_outline_only(self):
        import docx
        d = docx.Document()
        for i in range(30):
            d.add_heading('節%d' % i, level=1)
            d.add_paragraph('本文' * 400)
        src = self.folder/'long.docx'
        d.save(str(src))
        out = self.run_script('inspect_docx.py', src).stdout
        self.assertIn('§30 # 節29', out)
        self.assertNotIn('本文本文', out)
        self.assertIn('本文本文', self.run_script('inspect_docx.py', src, '--full').stdout)

    def test_dry_run_matches_real_run_and_covers_every_story(self):
        import re
        dry, out = self.replace(self.MANUAL, '--find', 'Alpha 2', '--replace', 'Beta 3', '--dry-run')
        self.assertEqual(dry.returncode, 0, dry.stderr)
        self.assertFalse(out.exists())
        real, out = self.replace(self.MANUAL, '--find', 'Alpha 2', '--replace', 'Beta 3')
        self.assertEqual(real.returncode, 0, real.stderr)
        line = '置換 7 件（本文 3、本文のテキストボックス 1、本文・他の変更履歴の挿入の中 1、脚注 1、ヘッダー 1）'
        self.assertIn(line, dry.stdout)
        self.assertIn(line, real.stdout)
        text = self.run_script('inspect_docx.py', out).stdout
        self.assertNotIn('Alpha 2', text)
        self.assertIn('型番\tBeta 3', text)
        # テキストボックスは Choice と Fallback の2通り保存されている。両方とも置換する
        with zipfile.ZipFile(out) as z:
            xml = z.read('word/document.xml').decode('utf-8')
        boxes = re.findall(r'<w:txbxContent>(.*?)</w:txbxContent>', xml)
        self.assertEqual(len(boxes), 2)
        self.assertTrue(all('Beta 3' in b and 'Alpha 2' not in b for b in boxes))

    def test_chained_finds_are_counted_like_the_real_run(self):
        dry, _ = self.replace(self.MANUAL, '--find', 'Alpha 2', '--replace', 'X', '--find', 'X', '--replace', 'Y',
                              '--dry-run')
        self.assertIn('連鎖', dry.stdout)
        self.assertIn('X → Y: 置換 7 件', dry.stdout)

    def test_match_across_tab_break_or_note_is_not_replaced(self):
        import docx
        d = docx.Document()
        p = d.add_paragraph('項目A')
        p.runs[0].add_tab()
        p.add_run('旧製品')
        p2 = d.add_paragraph('改行前')
        p2.runs[0].add_break()
        p2.add_run('改行後')
        src = self.folder/'tab.docx'
        d.save(str(src))
        for find in ('A旧', '前改'):
            result, out = self.replace(src, '--find', find, '--replace', 'X')
            self.assertEqual(result.returncode, 1, find)
            self.assertFalse(out.exists())
        result, out = self.replace(self.MANUAL, '--find', 'である。', '--replace', 'だ。')
        self.assertEqual(result.returncode, 1)
        self.assertIn('置換しなかった 1 件', result.stdout)
        self.assertFalse(out.exists())

    def test_track_splits_runs_and_keeps_format(self):
        import docx
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run('AB').bold = True
        p.add_run('CD')
        src = self.folder/'runs.docx'
        d.save(str(src))
        result, out = self.replace(src, '--find', 'BC', '--replace', 'X', '--track', '佐藤')
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn('変更履歴: 挿入 1・削除 1', result.stdout)
        body = docx.Document(str(out)).element.body
        W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        dels = [''.join(t.text for t in el.iter(W+'delText')) for el in body.iter(W+'del')]
        self.assertEqual(dels, ['BC'])
        ins = next(body.iter(W+'ins'))
        self.assertIsNotNone(ins.find('.//' + W + 'b'))            # 一致範囲の先頭（太字）の書式
        # Word は w:date を現地時刻として扱う（Word 2021 で確認）。現地時刻に Z を付けて書く
        import datetime
        self.assertEqual(ins.get(W+'date')[:13], datetime.datetime.now().strftime('%Y-%m-%dT%H'))
        text = self.run_script('inspect_docx.py', out).stdout
        self.assertIn('AXD', text)

    def test_track_stops_inside_other_insertion(self):
        result, out = self.replace(self.MANUAL, '--find', 'Alpha 2', '--replace', 'Beta 3', '--track', '佐藤')
        self.assertEqual(result.returncode, 1)
        self.assertIn('他の変更履歴の挿入の中', result.stdout)
        self.assertFalse(out.exists())
        result, out = self.replace(self.MANUAL, '--find', 'Alpha 2', '--replace', 'Beta 3', '--track', '佐藤',
                                   '--allow-missing')
        self.assertIn('変更履歴: 挿入 6・削除 6', result.stdout)
        import docx
        ids = [el.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id')
               for el in docx.Document(str(out)).element.body.iter()
               if isinstance(el.tag, str) and el.tag.rsplit('}', 1)[-1] in ('ins', 'del')]
        self.assertEqual(len(ids), len(set(ids)))
        summary = self.run_script('inspect_docx.py', out).stdout
        self.assertIn('変更履歴の挿入 7、変更履歴の削除 6', summary)

    def xml_doc(self, name, body_xml, setup=None):
        """本文の XML を差し込んだ文書を作る。"""
        import docx
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        d = docx.Document()
        body = d.element.body
        sect = body[-1]
        for el in parse_xml('<w:body %s %s>%s</w:body>' % (nsdecls('w'), nsdecls('r'), body_xml)):
            sect.addprevious(el)
        if setup:
            setup(d)
        path = self.folder/name
        d.save(str(path))
        return path

    def test_boundaries_from_review(self):
        # フィールドの境目・リンクの境目をまたぐ一致は置換しない
        src = self.xml_doc('field.docx',
                           '<w:p><w:r><w:t>A</w:t></w:r><w:fldSimple w:instr="PAGE"><w:r><w:t>1</w:t></w:r>'
                           '</w:fldSimple></w:p>'
                           '<w:p><w:r><w:t>A</w:t></w:r><w:hyperlink r:id="rIdX"><w:r><w:t>B</w:t></w:r>'
                           '</w:hyperlink></w:p>')
        for find in ('A1', 'AB'):
            result, out = self.replace(src, '--find', find, '--replace', 'X', '--track', 'R')
            self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
            self.assertIn('置換しなかった 1 件', result.stdout)
            self.assertFalse(out.exists())
        # フィールドの結果の中だけの一致は置換し、場所に出す
        result, _ = self.replace(src, '--find', '1', '--replace', '2', '--dry-run')
        self.assertIn('本文・フィールドの結果 1', result.stdout)
        # 空の w:t を挟んでも --track が落ちない
        src = self.xml_doc('empty.docx', '<w:p><w:r><w:t>A</w:t></w:r><w:r><w:t/></w:r><w:r><w:t>B</w:t></w:r></w:p>')
        result, _ = self.replace(src, '--find', 'AB', '--replace', 'X', '--track', 'R')
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('変更履歴: 挿入 1・削除 1', result.stdout)
        # 削除済みのテキストボックスの中は対象外
        box = ('<w:p><w:del w:id="5" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:pict><v:rect '
               'xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox><w:txbxContent><w:p><w:r><w:t>AB</w:t></w:r>'
               '</w:p></w:txbxContent></v:textbox></v:rect></w:pict></w:r></w:del></w:p>')
        result, _ = self.replace(self.xml_doc('delbox.docx', box), '--find', 'AB', '--replace', 'X')
        self.assertEqual(result.returncode, 1)
        self.assertIn('置換 0 件', result.stdout)

    def test_inspect_tables_headings_and_long_plain_text(self):
        cell = '<w:tc><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:tc>'
        nested = ('<w:tbl><w:tr><w:tc><w:p><w:r><w:t>OUTER</w:t></w:r></w:p><w:tbl><w:tr>' + cell % 'INNER' +
                  '</w:tr></w:tbl><w:p/></w:tc></w:tr></w:tbl>'
                  '<w:tbl><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid><w:tr>' + cell % 'H1' + cell % 'H2' +
                  '</w:tr><w:tr><w:trPr><w:gridBefore w:val="1"/></w:trPr>' + cell % 'VALUE' + '</w:tr></w:tbl>'
                  '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr></w:p>'
                  '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>本当の見出し</w:t></w:r></w:p>')
        src = self.xml_doc('tables.docx', nested)
        out = self.run_script('inspect_docx.py', src).stdout
        self.assertIn('| OUTER [表: INNER] |', out)
        self.assertIn('|  | VALUE |', out)
        self.assertIn('§1 # 本当の見出し', self.run_script('inspect_docx.py', src, '--outline').stdout)
        # 見出しの無い長い文書は先頭だけを出す
        import docx
        d = docx.Document()
        for _ in range(30):
            d.add_paragraph('本文' * 400)
        long = self.folder/'plain.docx'
        d.save(str(long))
        out = self.run_script('inspect_docx.py', long).stdout
        self.assertIn('見出しが無いので先頭', out)
        self.assertLess(len(out), 25000)

    def test_review_markup_helpers_use_local_time_and_fresh_ids(self):
        import re
        import docx
        md = (WORD/'references/review-markup.md').read_text(encoding='utf-8')
        g = {}
        exec(re.findall(r"```python\n(.*?)```", md, re.S)[1], g)
        import datetime
        self.assertTrue(g['NOW'].startswith(datetime.datetime.now().strftime('%Y-%m-%dT%H')))
        d = docx.Document(str(self.MANUAL))
        p = d.add_paragraph('消す文')
        g['tracked_insert'](d, p, '足す文', 'R')
        g['tracked_delete'](d, p.runs[0], 'R')
        out = self.folder/'rev.docx'
        d.save(str(out))
        summary = self.run_script('inspect_docx.py', out).stdout
        self.assertIn('変更履歴の挿入 2、変更履歴の削除 1', summary)
        ids = [el.get(g['qn']('w:id')) for el in d.element.body.iter() if isinstance(el.tag, str)
               and el.tag in (g['qn']('w:ins'), g['qn']('w:del'))]
        self.assertEqual(len(ids), len(set(ids)))   # 既存の 901 と重ならない


@unittest.skipUnless(HAS_PYPDF, 'pypdf が無い')
class PdfTests(TemporaryCase):
    FILES = PDF/'evals/files'
    SHORT_PDF = (b'%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n'
                 b'2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n'
                 b'3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 595 842]/Contents 4 0 R'
                 b'/Resources<</Font<</F1 5 0 R>>>>>>endobj\n'
                 b'4 0 obj<</Length 44>>stream\nBT /F1 12 Tf 72 720 Td (Short page) Tj ET\nendstream endobj\n'
                 b'5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n'
                 b'trailer<</Root 1 0 R>>\nstartxref\n0\n%%EOF\n')

    def run_script(self, script, *args):
        result = subprocess.run([sys.executable, str(PDF/'scripts'/script), *map(str, args)],
                                capture_output=True, text=True, encoding='utf-8',
                                env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        self.assertNotIn('Traceback', result.stderr)
        return result

    def meta(self, pdf):
        result = self.run_script('pdf_extract.py', 'meta', pdf)
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads(result.stdout)

    def pngs(self, pdf, dpi, count):
        result = self.run_script('pdf_to_png.py', pdf, '--outdir', self.folder, '--dpi', dpi)
        self.assertEqual(result.returncode, 0, result.stderr)
        paths = result.stdout.splitlines()
        self.assertEqual(len(paths), count)
        return paths

    def test_text_reports_image_only_page_in_mixed_document(self):
        # 4ページ中1ページだけ画像。全体の平均では文字数が足りてしまい、以前は終了コード0だった
        result = self.run_script('pdf_extract.py', 'text', self.FILES/'minutes.pdf')
        self.assertEqual(result.returncode, 1)
        self.assertIn('文字がほとんど取れないページ: 3（', result.stderr)
        self.assertIn('リリースは11月20日', result.stdout)

    def test_text_does_not_flag_short_page_without_image(self):
        # 本文の短いページ（Windows の Word で18字の文書）を、スキャンの疑いとして挙げていた
        garbage = self.file('broken.pdf', 'not a pdf')
        broken = self.run_script('pdf_extract.py', 'text', garbage)  # 壊れた PDF はトレースバックしていた
        self.assertEqual(broken.returncode, 1)
        self.assertIn('PDF を開けない', broken.stderr)
        self.assertEqual(self.run_script('pdf_ops.py', 'list-fields', garbage).returncode, 1)
        short = self.folder/'short.pdf'
        short.write_bytes(self.SHORT_PDF)
        result = self.run_script('pdf_extract.py', 'text', short)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('Short page', result.stdout)

    def test_bad_page_spec_is_usage_error(self):
        # 範囲外・逆順は以前、何もせず終了コード0（rotate は未変更の複製を書き出していた）
        src = self.FILES/'minutes.pdf'
        for spec in ('9', '3-1', 'a', ',', '1-', '0', '', '1 2'):
            for args in (('pdf_extract.py', 'text', src, '--pages', spec),
                         ('pdf_ops.py', 'rotate', src, self.folder/'r.pdf', '--degrees', '90', '--pages', spec),
                         ('pdf_ops.py', 'split', src, '--outdir', self.folder, '--ranges', spec),
                         ('pdf_ops.py', 'stamp', src, self.folder/'r.pdf',
                          '--stamp', self.FILES/'stamp_confidential.pdf', '--pages', spec),
                         ('pdf_to_png.py', src, '--outdir', self.folder, '--pages', spec)):
                self.assertEqual(self.run_script(*args).returncode, 2, (args[0], args[1], spec))
        self.assertEqual(list(self.folder.iterdir()), [])
        self.assertEqual(self.run_script('pdf_extract.py', 'text', src, '--pages', '1, 2').returncode, 0)

    def test_large_text_needs_out(self):
        big = self.folder/'big.pdf'
        from pypdf import PdfWriter
        w = PdfWriter()
        for _ in range(40):
            w.append(str(self.FILES/'minutes.pdf'))
        w.write(str(big))
        result = self.run_script('pdf_extract.py', 'text', big)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(result.stdout, '')
        out = self.folder/'big.txt'
        result = self.run_script('pdf_extract.py', 'text', big, '--out', out)
        self.assertEqual(result.stdout.strip(), str(out))
        self.assertIn('=== page 160 ===', out.read_text(encoding='utf-8'))

    def test_meta_shows_size_and_rotation_per_page(self):
        self.assertEqual(self.meta(self.FILES/'mixed_orientation.pdf')['page_layout'],
                         ['1: 210x297mm', '2: 297x210mm', '3: 210x297mm 回転90'])

    def test_encrypted_tables_is_not_traceback(self):
        enc = self.folder/'enc.pdf'
        r = self.run_script('pdf_ops.py', 'encrypt', self.FILES/'minutes.pdf', enc, '--password', 'pw')
        if r.returncode == 2:
            self.skipTest('cryptography が無い')
        self.assertEqual(self.run_script('pdf_extract.py', 'tables', enc).returncode, 1)

    def edited_form(self, edit):
        from pypdf import PdfWriter
        w = PdfWriter(clone_from=str(self.FILES/'application_form.pdf'))
        edit(w, {str(f.get_object()['/T']): f for f in w._root_object['/AcroForm']['/Fields']})
        out = self.folder/'form.pdf'
        w.write(str(out))
        return out

    def fill(self, values, form=None):
        data = self.file('v.json', json.dumps(values, ensure_ascii=False))
        out = self.folder/'filled.pdf'
        return self.run_script('pdf_ops.py', 'fill-form', form or self.FILES/'application_form.pdf',
                               out, '--data', data), out

    def test_fill_form_rejects_bad_values(self):
        # 以前は選択肢に無い "Z" も、テキスト欄への数値も、そのまま入れて終了コード0だった
        for values, message in (({'dept': 'Z'}, "dept: 'Z' は選択肢に無い"),
                                ({'employee_no': 1234}, 'employee_no: 値は文字列で書く')):
            result, out = self.fill(values)
            self.assertEqual(result.returncode, 1)
            self.assertIn(message, result.stderr)
            self.assertFalse(out.exists())
        self.assertEqual(self.fill({})[0].returncode, 2)

    def test_fill_form_checks_box_without_leading_slash(self):
        # 以前は "Yes" を入れると /V だけ変わり、チェックは付かないまま終了コード0だった
        result, out = self.fill({'agree': 'Yes', 'style': 'onsite', 'name': '山田花子'})
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('記入したページ: 1', result.stderr)
        self.assertIn('フィールドのフォントで表せない文字', result.stderr)
        from pypdf import PdfReader
        states = {}
        for a in PdfReader(str(out)).pages[0]['/Annots']:
            a = a.get_object()
            name = a.get('/T') or a['/Parent'].get_object()['/T']
            if a.get('/AS') not in (None, '/Off'):
                states[str(name)] = str(a['/AS'])
        self.assertEqual(states, {'agree': '/Yes', 'style': '/onsite'})

    def test_fill_form_refuses_same_leaf_name(self):
        # pypdf は末端の /T でも一致させるので、"name" を渡すと "group.name" にも書かれていた
        from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject

        def edit(w, fields):
            kid = fields['employee_no']
            kid.get_object()[NameObject('/T')] = TextStringObject('name')
            group = w._add_object(DictionaryObject({NameObject('/T'): TextStringObject('group'),
                                                    NameObject('/Kids'): ArrayObject([kid])}))
            kid.get_object()[NameObject('/Parent')] = group
            top = w._root_object['/AcroForm']['/Fields']
            top[list(top).index(kid)] = group
        form = self.edited_form(edit)
        result, out = self.fill({'name': '山田'}, form)
        self.assertEqual(result.returncode, 1)
        self.assertIn('group.name', result.stderr)
        result, out = self.fill({'group.name': 'A1234'}, form)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_fill_form_uses_export_value_of_option_pair(self):
        # /Opt が [保存値, 表示名] の組だと、以前は配列の文字列表現を選択肢として出していた
        from pypdf.generic import ArrayObject, NameObject, TextStringObject

        def edit(w, fields):
            fields['dept'].get_object()[NameObject('/Opt')] = ArrayObject(
                [ArrayObject([TextStringObject(v), TextStringObject(l)])
                 for v, l in (('dev', '開発'), ('qa', '品質保証'), ('sales', '営業'))])
        form = self.edited_form(edit)
        listing = self.run_script('pdf_ops.py', 'list-fields', form)
        dept = [f for f in json.loads(listing.stdout) if f['name'] == 'dept'][0]
        self.assertEqual((dept['options'], dept['labels']), (['dev', 'qa', 'sales'], ['開発', '品質保証', '営業']))
        result, _ = self.fill({'dept': '品質保証'}, form)
        self.assertEqual(result.returncode, 1)
        self.assertIn('options の値を使う（qa）', result.stderr)
        self.assertEqual(self.fill({'dept': 'qa'}, form)[0].returncode, 0)

    @unittest.skipUnless(HAS_PDFIUM, 'pypdfium2 か Pillow が無い')
    def test_png_draws_form_values(self):
        # 以前は init_forms を呼ばず、記入した値もフィールドの枠も画像に出なかった
        from PIL import Image, ImageChops
        result, out = self.fill({'agree': '/Yes'})
        self.assertEqual(result.returncode, 0, result.stderr)
        inside = (183, 842 - 682, 197, 842 - 668)  # チェックボックスの枠の内側
        filled = Image.open(self.pngs(out, '72', 1)[0]).convert('L').crop(inside)
        blank = Image.open(self.pngs(self.FILES/'application_form.pdf', '72', 1)[0]).convert('L').crop(inside)
        self.assertLess(filled.getextrema()[0], 100)  # 印が描かれている
        self.assertIsNotNone(ImageChops.difference(filled, blank).getbbox())

    def red_box(self, path):
        from PIL import Image
        im = Image.open(path).convert('RGB')
        red = [(x, y) for y in range(im.height) for x in range(im.width)
               if (lambda p: p[0] > 200 and p[1] < 230 and p[0] - p[1] > 30)(im.getpixel((x, y)))]
        xs, ys = [x for x, _ in red], [y for _, y in red]
        return im, (min(xs), min(ys), max(xs) + 1, max(ys) + 1)

    @unittest.skipUnless(HAS_PDFIUM, 'pypdfium2 か Pillow が無い')
    def test_stamp_is_centered_and_upright_on_every_page(self):
        # 以前は横向きのページで上に寄り、回転90のページで横倒しになった
        from PIL import ImageChops, ImageStat
        out = self.folder/'stamped.pdf'
        r = self.run_script('pdf_ops.py', 'stamp', self.FILES/'mixed_orientation.pdf', out,
                            '--stamp', self.FILES/'stamp_confidential.pdf')
        self.assertEqual(r.returncode, 0, r.stderr)
        crops = []
        for path in self.pngs(out, '30', 3):
            im, (x0, y0, x1, y1) = self.red_box(path)
            self.assertAlmostEqual((x0 + x1) / 2 / im.width, 0.5, delta=0.05, msg=path)
            self.assertAlmostEqual((y0 + y1) / 2 / im.height, 0.5, delta=0.08, msg=path)
            crops.append(im.crop((x0, y0, x1, y1)).convert('L'))
        # 縦のページの印を基準に、横・回転90のページの印が同じ向き（逆さでもない）か比べる
        ref = crops[0]
        for crop in crops[1:]:
            same = ImageChops.difference(ref, crop.resize(ref.size))
            flipped = ImageChops.difference(ref, crop.rotate(180).resize(ref.size))
            self.assertLess(ImageStat.Stat(same).sum[0], ImageStat.Stat(flipped).sum[0])

    @unittest.skipUnless(HAS_PDFIUM, 'pypdfium2 か Pillow が無い')
    def test_stamp_keeps_rotation_of_stamp_page(self):
        # 印のページ自体に /Rotate があると、以前は回転を無視して重ねていた
        rotated = self.folder/'stamp90.pdf'
        self.assertEqual(self.run_script('pdf_ops.py', 'rotate', self.FILES/'stamp_confidential.pdf', rotated,
                                         '--degrees', '90').returncode, 0)
        out = self.folder/'stamped.pdf'
        r = self.run_script('pdf_ops.py', 'stamp', self.FILES/'mixed_orientation.pdf', out,
                            '--stamp', rotated, '--pages', '1')
        self.assertEqual(r.returncode, 0, r.stderr)
        _, (x0, y0, x1, y1) = self.red_box(self.pngs(out, '30', 3)[0])
        self.assertGreater(y1 - y0, x1 - x0)  # 印単体の見た目どおり縦書き

    def test_to_pdf_failure_keeps_no_stale_success(self):
        # 以前は出力先に前回の PDF が残っていると、変換に失敗しても終了コード0でそのパスを返した
        bad = self.folder/'bad.docx'
        bad.write_bytes(b'PK\x03\x04broken')
        stale = self.folder/'bad.pdf'
        stale.write_text('stale')
        result = self.run_script('to_pdf.py', bad)
        if result.returncode == 3:
            self.skipTest('変換エンジンが無い')
        self.assertEqual(result.returncode, 1)
        self.assertEqual(stale.read_text(), 'stale')
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()), ['bad.docx', 'bad.pdf'])

    def test_to_pdf_decodes_cp932_output(self):
        # 日本語 Windows の PowerShell は cp932 で出力し、以前は読み取りスレッドが落ちて Traceback になった
        to_pdf = load('to_pdf', 'operate-pdf/scripts/to_pdf.py')
        r = to_pdf.run([sys.executable, '-c', 'import sys; sys.stdout.buffer.write("変換失敗".encode("cp932") + b"\\xff")'])
        self.assertEqual(r.returncode, 0)
        self.assertIsInstance(r.stdout, str)


class DrawioTests(TemporaryCase):
    def generate(self, spec):
        src = self.file('d.spec.json', json.dumps(spec, ensure_ascii=False))
        out = self.folder/'d.drawio'
        r = subprocess.run([sys.executable, str(DRAWIO/'scripts/layout_from_spec.py'), str(src), '-o', str(out)],
                           capture_output=True, text=True, encoding='utf-8', env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        import xml.etree.ElementTree as ET
        cells = {c.get('id'): c for c in ET.parse(out).getroot().iter('mxCell')}
        return out, cells

    @staticmethod
    def style(cell):
        return dict(p.partition('=')[::2] for p in cell.get('style', '').split(';') if p)

    @staticmethod
    def geo(cells, cid):
        g = cells[cid].find('mxGeometry')
        x, y, w, h = (float(g.get(k, 0)) for k in ('x', 'y', 'width', 'height'))
        parent = cells[cid].get('parent')
        if parent not in ('0', '1'):
            px, py, _, _ = DrawioTests.geo(cells, parent)
            x, y = x + px, y + py
        return x, y, w, h

    def ends(self, cells, eid):
        """線の出口・入口の座標と、経由点"""
        c, st = cells[eid], self.style(cells[eid])
        sx, sy, sw, sh = self.geo(cells, c.get('source'))
        tx, ty, tw, th = self.geo(cells, c.get('target'))
        pts = [(float(p.get('x')), float(p.get('y'))) for p in c.iter('mxPoint') if p.get('as') is None]
        return ((sx + sw * float(st['exitX']), sy + sh * float(st['exitY'])),
                (tx + tw * float(st['entryX']), ty + th * float(st['entryY'])), pts)

    def test_ports_line_up_with_waypoints(self):
        # 出入口の比率を3桁に丸めていたので、出入口と経由点が 0.4px ずれて線がわずかに斜めに描かれた
        _, cells = self.generate({'layout': 'free', 'nodes': [
            {'id': 'off', 'label': '停止', 'at': [0, 0]},
            {'id': 'idle', 'label': '待機', 'near': {'right_of': 'off', 'gap': 180}},
            {'id': 'run', 'label': '運転', 'near': {'right_of': 'idle', 'gap': 180}}],
            'edges': [{'from': 'off', 'to': 'idle'}, {'from': 'idle', 'to': 'run'},
                      {'from': 'run', 'to': 'idle', 'label': '停止指令'}, {'from': 'idle', 'to': 'off', 'label': '電源OFF'}]})
        for eid in ('e_run_idle', 'e_idle_off'):
            start, end, pts = self.ends(cells, eid)
            self.assertTrue(pts)
            for port, pt in ((start, pts[0]), (end, pts[-1])):
                self.assertLess(min(abs(port[0] - pt[0]), abs(port[1] - pt[1])), 0.01, (eid, port, pt))

    def test_halfwidth_label_does_not_widen_box_like_fullwidth(self):
        _, cells = self.generate({'nodes': [{'id': 'a', 'label': 'Customer Web Browser'},
                                            {'id': 'b', 'label': '顧客のウェブブラウザ画面'}]})
        self.assertLess(self.geo(cells, 'a')[2], 200)   # 以前は全角と同じ 14px/字で 304px
        self.assertGreaterEqual(self.geo(cells, 'b')[2], 14 * 12)

    def test_edge_into_box_inside_group_does_not_detour_around_group(self):
        # 枠の見出し帯全体を避けていたので、上の層から下の層の箱へ入る線が枠の横を大回りした
        _, cells = self.generate({'layout': 'free',
            'groups': [{'id': 'app', 'label': 'アプリケーション層', 'at': [0, 0]},
                       {'id': 'rte', 'label': 'RTE', 'near': {'below': 'app', 'gap': 60, 'align': 'start'}}],
            'nodes': [{'id': 's1', 'label': 'エンジン制御SWC', 'in': 'app', 'at': [0, 0]},
                      {'id': 's2', 'label': '診断SWC', 'in': 'app', 'near': {'right_of': 's1'}},
                      {'id': 's3', 'label': '通信SWC', 'in': 'app', 'near': {'right_of': 's2'}},
                      {'id': 'r', 'label': 'RTE', 'in': 'rte', 'at': [0, 0], 'size': [620, 50]}],
            'edges': [{'from': 's1', 'to': 'r'}, {'from': 's2', 'to': 'r'}, {'from': 's3', 'to': 'r'}]})
        for eid in ('e_s1_r', 'e_s3_r'):
            start, end, pts = self.ends(cells, eid)
            self.assertAlmostEqual(start[0], end[0], delta=0.01)   # まっすぐ下りる
            self.assertFalse(pts)
        gx, _, gw, _ = self.geo(cells, 'rte')
        for eid in ('e_s1_r', 'e_s2_r', 'e_s3_r'):
            for x, _ in self.ends(cells, eid)[2]:
                self.assertTrue(gx < x < gx + gw, eid)   # 枠の外を回らない

    def test_decision_branches_leave_from_different_vertices(self):
        _, cells = self.generate({'layout': 'grid', 'direction': 'TB', 'nodes': [
            {'id': 'in', 'label': '入力'}, {'id': 'chk', 'label': '形式は正しい?', 'shape': 'decision'},
            {'id': 'auth', 'label': '照会'}, {'id': 'log', 'label': 'エラーログ記録'}, {'id': 'lock', 'label': 'ロック'}],
            'edges': [{'from': 'in', 'to': 'chk'}, {'from': 'chk', 'to': 'auth', 'label': 'はい'},
                      {'from': 'chk', 'to': 'log', 'label': 'いいえ'}, {'from': 'auth', 'to': 'lock'},
                      {'from': 'lock', 'to': 'log'}, {'from': 'log', 'to': 'in'}]})
        sides = {(self.style(c)['exitX'], self.style(c)['exitY']) for c in cells.values()
                 if c.get('edge') == '1' and c.get('source') == 'chk'}
        self.assertEqual(len(sides), 2)

    def test_edges_of_different_kinds_do_not_share_a_port(self):
        # 汎化と依存を同じ点から出すと、1本の線が箱の横を通り抜けているように見えた
        _, cells = self.generate({'layout': 'free', 'nodes': [
            {'id': 'base', 'label': 'SensorBase', 'shape': 'class', 'methods': ['+ read(): int32_t'], 'at': [200, 0]},
            {'id': 'temp', 'label': 'TempSensor', 'shape': 'class', 'near': {'below': 'base', 'gap': 100, 'align': 'start'}},
            {'id': 'pres', 'label': 'PressureSensor', 'shape': 'class', 'near': {'right_of': 'temp', 'gap': 100}},
            {'id': 'adc', 'label': 'AdcDriver', 'shape': 'class', 'near': {'below': 'temp', 'gap': 100}}],
            'edges': [{'from': 'temp', 'to': 'base', 'kind': 'inherit'}, {'from': 'pres', 'to': 'base', 'kind': 'inherit'},
                      {'from': 'temp', 'to': 'adc', 'kind': 'depend'}, {'from': 'pres', 'to': 'adc', 'kind': 'depend'}]})
        a, b = self.ends(cells, 'e_pres_base')[0], self.ends(cells, 'e_pres_adc')[0]
        self.assertGreater(abs(a[0] - b[0]) + abs(a[1] - b[1]), 5)

    def test_export_rejects_page_beyond_page_count(self):
        # draw.io はページ数を超える番号でも黙って1ページ目を書き出す
        out, _ = self.generate({'nodes': [{'id': 'a', 'label': 'A'}]})
        r = subprocess.run([sys.executable, str(DRAWIO/'scripts/export_drawio.py'), str(out), '--page', '3',
                            '--exe', sys.executable], capture_output=True, text=True, encoding='utf-8',
                           env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        self.assertEqual(r.returncode, 2, r.stdout)

    def test_export_launches_on_linux_as_root_without_display(self):
        # root で --no-sandbox が無いと draw.io が起動直後に落ち、終了コード 4（権限の制限）で止まっていた
        export = load('export_drawio', 'drawio-diagram/scripts/export_drawio.py')
        env = {k: v for k, v in os.environ.items() if k not in ('DISPLAY', 'WAYLAND_DISPLAY')}
        with patch.object(export.platform, 'system', return_value='Linux'), \
                patch.object(export.os, 'geteuid', return_value=0, create=True), \
                patch.dict(export.os.environ, env, clear=True), \
                patch.object(export.shutil, 'which', return_value='/usr/bin/xvfb-run'):
            pre, post = export.launcher('/opt/drawio/drawio')
        self.assertEqual(pre, ['xvfb-run', '-a', '/opt/drawio/drawio'])
        self.assertEqual(post, ['--no-sandbox'])
        with patch.object(export.platform, 'system', return_value='Windows'):
            self.assertEqual(export.launcher('draw.io.exe'), (['draw.io.exe'], []))

    def test_header_check_flags_only_lines_over_the_title(self):
        export = load('export_drawio', 'drawio-diagram/scripts/export_drawio.py')
        exe = export.find_exe(None)
        if not exe:
            self.skipTest('draw.io が無い')
        doc = self.file('h.drawio', """<mxfile><diagram id="p" name="p"><mxGraphModel><root>
<mxCell id="0"/><mxCell id="1" parent="0"/>
<mxCell id="top1" value="A" style="rounded=0;html=1;" vertex="1" parent="1"><mxGeometry x="40" y="0" width="80" height="40" as="geometry"/></mxCell>
<mxCell id="top2" value="B" style="rounded=0;html=1;" vertex="1" parent="1"><mxGeometry x="200" y="0" width="80" height="40" as="geometry"/></mxCell>
<mxCell id="g" value="見出し" style="swimlane;startSize=30;html=1;" vertex="1" parent="1"><mxGeometry x="0" y="100" width="480" height="160" as="geometry"/></mxCell>
<mxCell id="in1" value="C" style="rounded=0;html=1;" vertex="1" parent="g"><mxGeometry x="40" y="80" width="80" height="40" as="geometry"/></mxCell>
<mxCell id="in2" value="D" style="rounded=0;html=1;" vertex="1" parent="g"><mxGeometry x="200" y="80" width="80" height="40" as="geometry"/></mxCell>
<mxCell id="ok" style="edgeStyle=none;html=1;exitX=0.5;exitY=1;entryX=0.5;entryY=0;" edge="1" parent="1" source="top1" target="in1"><mxGeometry relative="1" as="geometry"/></mxCell>
<mxCell id="ng" style="edgeStyle=none;html=1;exitX=0.5;exitY=1;entryX=0.5;entryY=0;" edge="1" parent="1" source="top2" target="in2"><mxGeometry relative="1" as="geometry"/></mxCell>
</root></mxGraphModel></diagram></mxfile>""")
        svg = self.folder/'h.svg'
        r = subprocess.run([sys.executable, str(DRAWIO/'scripts/export_drawio.py'), str(doc), '-f', 'svg', '-o', str(svg)],
                           capture_output=True, text=True, encoding='utf-8', env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        if r.returncode == 4:
            self.skipTest('draw.io が起動できない環境')
        self.assertEqual(r.returncode, 0, r.stdout)
        r = subprocess.run([sys.executable, str(DRAWIO/'scripts/check_layout.py'), str(doc), str(svg)],
                           capture_output=True, text=True, encoding='utf-8', env=dict(os.environ, PYTHONIOENCODING='utf-8'))
        self.assertIn('HEADER   線「ng」', r.stdout)
        self.assertNotIn('線「ok」', r.stdout)


if __name__ == '__main__':
    unittest.main(verbosity=2)
