#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Word 文書（.docx）の構造と本文を Markdown 風に出す。

    python inspect_docx.py <docx> [--outline] [--out FILE]

出力:
    - 冒頭: 用紙・向き・余白（セクションごと）、使われている段落スタイル、表・画像・コメント・
      変更履歴（挿入・削除）・フィールド（目次・ページ番号など）の数、ヘッダー・フッターの文字
    - 本文: 段落と表を文書の順に。見出しスタイルは #、箇条書きスタイルは -、
      それ以外は [スタイル名] を前に付ける（標準スタイルは付けない）。表は Markdown の表
    --outline  見出しだけを出す（長い文書の目次確認用）
    --out      本文をファイルに書き、標準出力にはパスだけを出す

終了コード:
    0  出力した
    1  開けない（壊れている、.docx ではない）
    2  引数の指定ミス、入力が存在しない、または python-docx が無い

依存: python-docx（pip install python-docx）
"""

import argparse
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

HEADING_RE = re.compile(r"^(Heading|見出し)\s*(\d)$")
TITLE_STYLES = {"Title", "表題"}
LIST_HINTS = ("List", "リスト", "箇条書き", "段落番号")
NORMAL_STYLES = {"Normal", "標準", "Body Text", "本文"}


def heading_level(style_name):
    if style_name in TITLE_STYLES:
        return 1
    m = HEADING_RE.match(style_name or "")
    return int(m.group(2)) if m else 0


def para_line(p):
    from docx.oxml.ns import qn
    text = p.text
    style = p.style.name if p.style is not None else ""
    level = heading_level(style)
    if level:
        return "#" * level + " " + text
    numbered = p._p.find(".//" + qn("w:numPr")) is not None
    if numbered or any(h in style for h in LIST_HINTS):
        return "- " + text
    if not text.strip():
        return ""
    if style and style not in NORMAL_STYLES:
        return "[%s] %s" % (style, text)
    return text


def table_lines(table):
    rows = []
    for r in table.rows:
        rows.append([c.text.replace("\n", " ").replace("|", "\\|") for c in r.cells])
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return out


def summary(doc, path):
    from docx.oxml.ns import qn
    lines = ["# %s" % path.name, ""]
    for i, s in enumerate(doc.sections, 1):
        if s.page_width is None:
            continue
        lines.append("- セクション%d: %.0f×%.0f mm（%s）、余白 上%.0f 下%.0f 左%.0f 右%.0f mm" % (
            i, s.page_width.mm, s.page_height.mm,
            "横" if s.page_width > s.page_height else "縦",
            s.top_margin.mm, s.bottom_margin.mm, s.left_margin.mm, s.right_margin.mm))
        for label, part in (("ヘッダー", s.header), ("フッター", s.footer)):
            if not part.is_linked_to_previous or i == 1:
                text = " / ".join(p.text for p in part.paragraphs if p.text.strip())
                if text:
                    lines.append("  - %s: %s" % (label, text))
    styles = Counter(p.style.name for p in doc.paragraphs if p.style is not None and p.text.strip())
    lines.append("- 段落スタイル: " + ", ".join("%s(%d)" % kv for kv in styles.most_common()))
    body = doc.element.body
    # フィールドはページ番号のようにフッターにあることが多いので、ヘッダー・フッターも見る
    roots = [body] + [part._element for s in doc.sections for part in (s.header, s.footer)
                      if not part.is_linked_to_previous]
    fields = [f.get(qn("w:instr")) for r in roots for f in r.iter(qn("w:fldSimple"))]
    fields += [t.text for r in roots for t in r.iter(qn("w:instrText")) if t.text and t.text.strip()]
    counts = [
        ("表", len(doc.tables)),
        ("画像", len(doc.inline_shapes)),
        ("変更履歴の挿入", sum(1 for _ in body.iter(qn("w:ins")))),
        ("変更履歴の削除", sum(1 for _ in body.iter(qn("w:del")))),
    ]
    with zipfile.ZipFile(path) as z:
        if "word/comments.xml" in z.namelist():
            counts.append(("コメント", z.read("word/comments.xml").count(b"<w:comment ")))
    lines.append("- " + "、".join("%s %d" % kv for kv in counts))
    kinds = Counter(f.strip().split()[0] for f in fields if f and f.strip())
    lines.append("- フィールド: " + (", ".join("%s(%d)" % kv for kv in kinds.most_common()) or "なし"))
    lines.append("")
    return lines


def body_lines(doc, outline):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    out = []
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            line = para_line(block)
            if outline and not line.startswith("#"):
                continue
            if line or (out and out[-1] != ""):
                out.append(line)
        elif isinstance(block, Table) and not outline:
            out.append("")
            out += table_lines(block)
            out.append("")
    return out


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Word 文書の構造と本文を出す")
    ap.add_argument("docx")
    ap.add_argument("--outline", action="store_true", help="見出しだけ出す")
    ap.add_argument("--out", help="本文をこのファイルに書く")
    args = ap.parse_args()

    try:
        import docx
    except ImportError:
        print("python-docx が必要: pip install python-docx", file=sys.stderr)
        return 2
    path = Path(args.docx)
    if not path.is_file():
        print("入力が存在しない: %s" % path, file=sys.stderr)
        return 2
    if path.suffix.lower() == ".doc":
        print(".doc は読めない。Word か LibreOffice で .docx に変換してから扱う", file=sys.stderr)
        return 2
    try:
        doc = docx.Document(str(path))
    except Exception as e:  # 壊れた zip・docx 以外の形式で python-docx が投げる例外はまちまち
        print("開けない: %s (%s)" % (path, e), file=sys.stderr)
        return 1

    head = summary(doc, path)
    body = body_lines(doc, args.outline)
    if args.out:
        Path(args.out).write_text("\n".join(head + body), encoding="utf-8")
        print(args.out)
    else:
        print("\n".join(head + body))
    return 0


if __name__ == "__main__":
    sys.exit(main())
