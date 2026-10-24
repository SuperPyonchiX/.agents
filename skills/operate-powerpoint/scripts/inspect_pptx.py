#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PowerPoint（.pptx / .potx）の構造・文字・配置を出し、配置と文字の問題を検査する。

    python inspect_pptx.py <pptx> [--layouts] [--slides 1-3,5] [--outline] [--full] [--check]

出力（Markdown）:
    - 冒頭: スライドの大きさ（mm と縦横比）、スライド枚数、非表示スライド
    - --layouts: テンプレートのレイアウト一覧と、各レイアウトのプレースホルダ（idx・種類・位置）
    - スライドごと: レイアウト名、図形（名前・種類・プレースホルダ idx・位置と大きさ mm・文字）、ノート。
      グループの中の図形は字下げして出し、位置はスライド上の位置に直す。表はセルの文字を行ごとに出す
    - --outline: スライドごとに1行（番号・レイアウト・タイトル・図形数）だけ出す。枚数の多い資料の下見用
    - --full: 文字・ノートを省略せずに出す（既定は図形ごと60字、ノート120字で切る）
    --check: 次を「検査」節に出す
        問題（終了コード1）: スライド外へのはみ出し、図形どうしの重なり、空のプレースホルダ
        注意（終了コードに影響しない）: 文字のあふれの疑い（推定）、18pt 未満の文字

終了コード:
    0  出力した（--check では問題なし。注意だけなら0）
    1  開けない、または --check で問題が見つかった
    2  引数の指定ミス、入力が存在しない、.ppt、または python-pptx が無い

依存: python-pptx（pip install python-pptx）
"""

import argparse
import io
import math
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

EMU_PER_MM = 36000
EMU_PER_PT = 12700
TOLERANCE_MM = 1.0        # これ以下のはみ出し・重なりは無視する
MIN_PT = 18.0             # SKILL.md の「本文 18pt 以上」
DEFAULT_INSET_LR = 91440  # bodyPr の既定の内側余白（左右 0.1in、上下 0.05in）
DEFAULT_INSET_TB = 45720
LINE_HEIGHT = 1.2         # 行送り（フォントサイズに対する倍率）の概算
BACKGROUND_RATIO = 0.9    # スライドの9割以上を覆う画像は背景とみなし、重なりを見ない
# 本文として出さないプレースホルダ（日付・フッター・スライド番号）
FOOTER_TYPES = {"dt", "ftr", "sldNum", "hdr"}

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


class UsageError(Exception):
    pass


def mm(v):
    return (v or 0) / EMU_PER_MM


def parse_pages(spec, total):
    if not spec:
        return list(range(total))
    pages = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        a, _, b = part.partition("-")
        if not a.strip().isdigit() or (b and not b.strip().isdigit()):
            raise UsageError("--slides の書式が違う: %s（例: 1-3,5）" % spec)
        start, end = int(a), int(b or a)
        if start > end or start < 1 or end > total:
            raise UsageError("--slides %s は範囲外（スライドは 1〜%d）" % (part, total))
        pages.extend(range(start - 1, end))
    return sorted(set(pages))


def open_presentation(path):
    """.potx は python-pptx がそのままでは開かないので、種類だけ .pptx に読み替えて開く。"""
    from pptx import Presentation
    if path.suffix.lower() not in (".potx", ".potm"):
        return Presentation(str(path))
    buf = io.BytesIO()
    with zipfile.ZipFile(str(path)) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb"presentationml\.template\.main\+xml|ms-powerpoint\.template\.macroEnabled\.main\+xml",
                              b"presentationml.presentation.main+xml", data)
            dst.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


# ---------------------------------------------------------------- 図形の列挙

def walk(shapes, transform=None, depth=0):
    """(図形, スライド上の枠 mm, 深さ) を返す。グループの中は子座標をスライド座標に直す。"""
    for s in shapes:
        b = raw_box(s)
        if b is not None and transform:
            b = transform(b)
        yield s, b, depth
        if s.shape_type is not None and str(s.shape_type).startswith("GROUP"):
            xfrm = s._element.grpSpPr.find(A + "xfrm")
            t = group_transform(xfrm, transform)
            yield from walk(s.shapes, t, depth + 1)


def raw_box(shape):
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if left is None or width is None:
        return None
    return (left, top, left + width, top + height)


def group_transform(xfrm, outer):
    if xfrm is None:
        return outer
    def get(tag, attr):
        el = xfrm.find(A + tag)
        return int(el.get(attr)) if el is not None else None
    ox, oy, ex, ey = get("off", "x"), get("off", "y"), get("ext", "cx"), get("ext", "cy")
    cx, cy, cex, cey = get("chOff", "x"), get("chOff", "y"), get("chExt", "cx"), get("chExt", "cy")
    if None in (ox, oy, ex, ey):
        return outer
    cx, cy = (ox, oy) if cx is None else (cx, cy)
    sx = ex / cex if cex else 1.0
    sy = ey / cey if cey else 1.0

    def t(b):
        x1, y1, x2, y2 = b
        r = (ox + (x1 - cx) * sx, oy + (y1 - cy) * sy, ox + (x2 - cx) * sx, oy + (y2 - cy) * sy)
        return outer(r) if outer else r
    return t


def to_mm(b):
    return tuple(mm(v) for v in b) if b else None


def is_group(s):
    return s.shape_type is not None and str(s.shape_type).startswith("GROUP")


def ph_type(shape):
    if not shape.is_placeholder:
        return None
    ph = shape._element.find(".//" + P + "ph")
    return ph.get("type", "body") if ph is not None else "body"


def shape_kind(shape):
    if shape.is_placeholder:
        return "ph[%d] %s" % (shape.placeholder_format.idx,
                              str(shape.placeholder_format.type).split(".")[-1].split(" ")[0])
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "グラフ"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "表 %d×%d" % (len(shape.table.rows), len(shape.table.columns))
    return str(shape.shape_type).split(".")[-1].split(" ")[0] if shape.shape_type else "図形"


def clip(text, limit):
    return text if limit is None or len(text) <= limit else text[: limit - 1] + "…"


def shape_text(shape, limit=60):
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [" | ".join(c.text.replace("\n", " ") for c in r.cells) for r in shape.table.rows]
        return clip(" / ".join(r for r in rows if r.strip(" |")), limit)
    if not shape.has_text_frame:
        return ""
    return clip(" / ".join(p.text for p in shape.text_frame.paragraphs if p.text.strip()), limit)


def has_text(shape):
    return bool(shape_text(shape, None))


# ---------------------------------------------------------------- 文字の大きさ

def level_style(lst, level):
    if lst is None:
        return None
    return lst.find(A + "lvl%dpPr" % (level + 1))


def size_in(lvl):
    if lvl is None:
        return None
    d = lvl.find(A + "defRPr")
    return int(d.get("sz")) / 100 if d is not None and d.get("sz") else None


class Styles:
    """プレースホルダの継承（スライド→レイアウト→マスター→txStyles→既定）をたどって文字の大きさを決める。"""

    def __init__(self, prs):
        el = prs.part._element
        self.default = el.find(P + "defaultTextStyle")

    def chain(self, shape):
        """文字の大きさを探す lstStyle の候補を、近い順に返す。"""
        out = []
        txbody = shape._element.find(P + "txBody")
        if txbody is not None:
            out.append(txbody.find(A + "lstStyle"))
        if shape.is_placeholder:
            base = shape
            for _ in range(2):  # レイアウト → マスター
                try:
                    base = base._base_placeholder
                except Exception:
                    base = None
                if base is None:
                    break
                tb = base._element.find(P + "txBody")
                if tb is not None:
                    out.append(tb.find(A + "lstStyle"))
            master = shape.part.slide_layout.slide_master if hasattr(shape.part, "slide_layout") else None
            if master is not None:
                styles = master._element.find(P + "txStyles")
                if styles is not None:
                    t = ph_type(shape)
                    tag = "titleStyle" if t in ("title", "ctrTitle") else \
                        "otherStyle" if t in FOOTER_TYPES else "bodyStyle"
                    out.append(styles.find(P + tag))
        out.append(self.default)
        return out

    def paragraph_sizes(self, shape, text_frame):
        """段落ごとに (段落の文字, 最大の文字の大きさ pt, 最小の大きさ pt) を返す。"""
        chain = self.chain(shape)
        result = []
        for p in text_frame.paragraphs:
            level = p.level
            inherited = next((s for s in (size_in(level_style(c, level)) for c in chain) if s), 18.0)
            sizes = [r.font.size.pt if r.font.size is not None else inherited for r in p.runs] or [inherited]
            result.append((p.text, max(sizes), min(sizes)))
        return result


def font_scale(text_frame):
    """normAutofit の fontScale（PowerPoint が縮小して保存した倍率）。無ければ 1.0。"""
    body = text_frame._txBody.find(A + "bodyPr")
    fit = body.find(A + "normAutofit") if body is not None else None
    if fit is not None and fit.get("fontScale"):
        return int(fit.get("fontScale")) / 100000
    return 1.0


def text_width_em(text):
    return sum(1.0 if unicodedata.east_asian_width(c) in ("W", "F", "A") else 0.55 for c in text)


def overflow_estimate(shape, b, styles):
    """枠に対して文字が多すぎるかを概算する。(必要な高さ mm, 使える高さ mm) か、判定できなければ None。"""
    tf = shape.text_frame
    body = tf._txBody.find(A + "bodyPr")
    get = (lambda k, d: int(body.get(k)) if body is not None and body.get(k) else d)
    avail_w = (b[2] - b[0]) * EMU_PER_MM - get("lIns", DEFAULT_INSET_LR) - get("rIns", DEFAULT_INSET_LR)
    avail_h = (b[3] - b[1]) * EMU_PER_MM - get("tIns", DEFAULT_INSET_TB) - get("bIns", DEFAULT_INSET_TB)
    if avail_w <= 0 or avail_h <= 0:
        return None
    wrap = body is None or body.get("wrap") != "none"
    scale = font_scale(tf)
    need = 0.0
    widest = 0.0
    for text, big, _ in styles.paragraph_sizes(shape, tf):
        size = big * scale * EMU_PER_PT
        lines = 0
        for line in (text.split("\v") if text else [""]):
            w = text_width_em(line) * size
            widest = max(widest, w)
            lines += max(1, math.ceil(w / avail_w)) if wrap else 1
        need += lines * size * LINE_HEIGHT
    if not wrap:
        # 折り返さない枠は横にはみ出す。高さではなく幅で判定する
        return ("幅", widest / EMU_PER_MM, avail_w / EMU_PER_MM) if widest > avail_w * 1.05 else None
    return ("高さ", need / EMU_PER_MM, avail_h / EMU_PER_MM) if need > avail_h * 1.05 else None


# ---------------------------------------------------------------- 検査

def check_slide(slide, width, height, styles):
    problems, notes = [], []
    targets = []
    for s, raw, _ in walk(slide.shapes):
        b = to_mm(raw)
        if b is None or is_group(s):
            continue
        x1, y1, x2, y2 = b
        if x1 < -TOLERANCE_MM or y1 < -TOLERANCE_MM or x2 > width + TOLERANCE_MM or y2 > height + TOLERANCE_MM:
            problems.append("「%s」がスライドの外にはみ出している（%.0f,%.0f〜%.0f,%.0f mm）" % (s.name, x1, y1, x2, y2))
        if s.is_placeholder and s._element.tag == P + "sp" and ph_type(s) not in FOOTER_TYPES and not has_text(s):
            problems.append("「%s」が空のプレースホルダ（画像には写らない。文字を入れるか消す）" % s.name)
        area = (x2 - x1) * (y2 - y1)
        is_picture = s._element.tag == P + "pic"
        is_frame = getattr(s, "has_table", False) and s.has_table or getattr(s, "has_chart", False) and s.has_chart
        if has_text(s) or is_frame or (is_picture and area < width * height * BACKGROUND_RATIO):
            targets.append((s, b))
        if s.has_text_frame and has_text(s):
            scale = font_scale(s.text_frame)
            small = [sm * scale for _, _, sm in styles.paragraph_sizes(s, s.text_frame)]
            if small and min(small) < MIN_PT - 0.01 and ph_type(s) not in FOOTER_TYPES:
                notes.append("「%s」に %.0fpt の文字がある（本文は %.0fpt 以上）" % (s.name, min(small), MIN_PT))
            est = overflow_estimate(s, b, styles)
            if est:
                notes.append("「%s」の文字があふれている疑い（推定: 必要な%s %.0f mm／枠 %.0f mm）"
                             % (s.name, est[0], est[1], est[2]))
    for i in range(len(targets)):
        for j in range(i + 1, len(targets)):
            (sa, a), (sb, b) = targets[i], targets[j]
            if not (has_text(sa) or has_text(sb)):
                continue  # 文字の無い図形どうし（画像と画像など）は見ない
            w = min(a[2], b[2]) - max(a[0], b[0])
            h = min(a[3], b[3]) - max(a[1], b[1])
            if w > TOLERANCE_MM and h > TOLERANCE_MM:
                problems.append("「%s」と「%s」が重なっている（%.0f×%.0f mm）" % (sa.name, sb.name, w, h))
    return problems, notes


# ---------------------------------------------------------------- 出力

def slide_title(slide):
    for s in slide.placeholders:
        if ph_type(s) in ("title", "ctrTitle") and s.has_text_frame and s.text_frame.text.strip():
            return s.text_frame.text.strip().replace("\n", " ").replace("\v", " ")
    return "（タイトルなし）"


def is_hidden(slide):
    return slide._element.get("show") in ("0", "false")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="pptx の構造と配置を出す")
    ap.add_argument("pptx")
    ap.add_argument("--layouts", action="store_true", help="レイアウト一覧も出す")
    ap.add_argument("--slides", help="例: 1-3,5（1始まり）")
    ap.add_argument("--outline", action="store_true", help="スライドごとに1行だけ出す")
    ap.add_argument("--full", action="store_true", help="文字とノートを省略しない")
    ap.add_argument("--check", action="store_true", help="はみ出し・重なり・空のプレースホルダ・あふれの疑い・小さい文字を検査する")
    args = ap.parse_args()

    try:
        import pptx  # noqa: F401
    except ImportError:
        print("python-pptx が必要: pip install python-pptx", file=sys.stderr)
        return 2
    path = Path(args.pptx)
    if not path.is_file():
        print("入力が存在しない: %s" % path, file=sys.stderr)
        return 2
    if path.suffix.lower() == ".ppt":
        print(".ppt は読めない。PowerPoint か LibreOffice で .pptx に変換してから扱う", file=sys.stderr)
        return 2
    try:
        prs = open_presentation(path)
    except Exception as e:  # 壊れた zip・形式違いで python-pptx が投げる例外はまちまち
        print("開けない: %s (%s)" % (path, e), file=sys.stderr)
        return 1
    try:
        pages = parse_pages(args.slides, len(prs.slides))
    except UsageError as e:
        print(e, file=sys.stderr)
        return 2

    limit = None if args.full else 60
    notes_limit = None if args.full else 120
    width, height = mm(prs.slide_width), mm(prs.slide_height)
    ratio = width / height if height else 0
    label = "16:9" if abs(ratio - 16 / 9) < 0.02 else "4:3" if abs(ratio - 4 / 3) < 0.02 else "%.2f:1" % ratio
    lines = ["# %s" % path.name, "",
             "- スライドの大きさ: %.0f × %.0f mm（%s）" % (width, height, label),
             "- スライド枚数: %d" % len(prs.slides)]
    hidden = [i + 1 for i, s in enumerate(prs.slides) if is_hidden(s)]
    if hidden:
        lines.append("- 非表示スライド: %s（PDF・画像には出ない。画像の番号は表示されるスライドだけで数える）"
                     % ", ".join(map(str, hidden)))
    lines.append("")

    if args.layouts:
        lines.append("## レイアウト")
        for m, master in enumerate(prs.slide_masters, 1):
            for layout in master.slide_layouts:
                lines.append("- マスター%d「%s」" % (m, layout.name))
                for ph in layout.placeholders:
                    if ph_type(ph) in FOOTER_TYPES:
                        continue
                    b = to_mm(raw_box(ph))
                    pos = "%.0f,%.0f %.0f×%.0f mm" % (b[0], b[1], b[2] - b[0], b[3] - b[1]) if b else "位置なし"
                    lines.append("  - %s 「%s」 %s" % (shape_kind(ph), ph.name, pos))
        lines.append("")

    styles = Styles(prs)
    all_problems, all_notes = [], []
    page_no = 0
    shown = {i: None for i in range(len(prs.slides))}
    for i, s in enumerate(prs.slides):
        if not is_hidden(s):
            page_no += 1
            shown[i] = page_no
    for i in pages:
        slide = prs.slides[i]
        tag = "非表示" if shown[i] is None else ("画像 p%d" % shown[i]) if hidden else ""
        head = "スライド %d（%s%s）" % (i + 1, slide.slide_layout.name, "・" + tag if tag else "")
        if args.outline:
            count = sum(1 for _ in walk(slide.shapes))
            lines.append("- %s: %s［図形 %d］" % (head, clip(slide_title(slide), limit), count))
        else:
            lines.append("## " + head)
            for s, raw, depth in walk(slide.shapes):
                b = to_mm(raw)
                pos = "%.0f,%.0f %.0f×%.0f" % (b[0], b[1], b[2] - b[0], b[3] - b[1]) if b else "-"
                text = shape_text(s, limit)
                lines.append("%s- %s 「%s」 %s mm%s" % ("  " * depth, shape_kind(s), s.name, pos,
                                                       ("： " + text) if text else ""))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append("- ノート： " + clip(notes.replace("\n", " / "), notes_limit))
            lines.append("")
        if args.check:
            problems, notes = check_slide(slide, width, height, styles)
            all_problems += ["スライド %d: %s" % (i + 1, p) for p in problems]
            all_notes += ["スライド %d: %s" % (i + 1, n) for n in notes]

    if args.check:
        if args.outline:
            lines.append("")
        lines.append("## 検査")
        lines += ["- " + p for p in all_problems] or ["- 問題なし（はみ出し・重なり・空のプレースホルダ）"]
        if all_notes:
            lines += ["", "### 注意（推定を含む。画像で確かめる）"] + ["- " + n for n in all_notes]
    print("\n".join(lines))
    return 1 if args.check and all_problems else 0


if __name__ == "__main__":
    sys.exit(main())
