#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""編集前後の Excel ブックを突き合わせ、変わったところを列挙する。

    python diff_xlsx.py <before.xlsx> <after.xlsx> [--limit 100]

既存ブックを編集したあと、意図したセル以外を変えていないかを確かめるために使う。
比べるもの: シートの集合と順序・表示状態、名前付き範囲、各シートの使用範囲、結合セル、ウィンドウ枠の固定、
印刷範囲、入力規則・条件付き書式の数、列幅・非表示列、行の高さ・非表示行、
セルの値（数式は式の文字列で比べる）、セルの書式（表示形式・フォント・塗りつぶし・罫線・配置）。
セルは実在するものだけを走査するので、書式だけのセルで使用範囲が広がったシートでも時間がかからない。

出力（Markdown）: 構造の差分と、変わったセルの範囲・件数・一覧（シートごと先頭 --limit 件）。

終了コード:
    0  差分なし
    1  差分あり（エラーではない。一覧が意図どおりかを確かめる）
    2  引数の指定ミス、入力が存在しない、または openpyxl が無い

依存: openpyxl（pip install openpyxl）
"""

import argparse
import sys
from pathlib import Path


def color(c):
    if c is None:
        return None
    return c.rgb if c.type == "rgb" else "%s:%s" % (c.type, c.value)


def style_key(c):
    f, a, b = c.font, c.alignment, c.border
    fill = color(c.fill.fgColor) if c.fill.fill_type else None
    return (c.number_format,
            (f.name, f.sz, f.b, f.i, f.u, f.strike, color(f.color)),
            (c.fill.fill_type, fill),
            tuple(getattr(getattr(b, side), "style", None) for side in ("left", "right", "top", "bottom")),
            (a.horizontal, a.vertical, a.wrap_text, a.indent),
            c.protection.locked)


def style_names(ka, kb):
    labels = ("表示形式", "フォント", "塗りつぶし", "罫線", "配置", "保護")
    return [label for label, x, y in zip(labels, ka, kb) if x != y]


def show(v):
    if v is None:
        return "（空）"
    v = getattr(v, "text", v)  # 配列数式は式の文字列で見せる
    s = str(v).replace("\n", " ").replace("|", "\\|")
    return s if len(s) <= 40 else s[:39] + "…"


def col_dims(ws):
    return {k: (d.width, d.hidden) for k, d in ws.column_dimensions.items() if d.customWidth or d.hidden}


def row_dims(ws):
    return {k: (d.ht, d.hidden) for k, d in ws.row_dimensions.items() if d.customHeight or d.hidden}


def value_key(v):
    # 配列数式（ArrayFormula）は等値比較を持たないので、範囲と式の文字列で比べる
    return (v.ref, v.text) if type(v).__name__ == "ArrayFormula" else v


def fmt_dim(d):
    if d is None:
        return "既定"
    size, hidden = d
    parts = (["%g" % size] if size is not None else []) + (["非表示"] if hidden else [])
    return "・".join(parts) or "既定"


def diff_dims(label, da, db):
    keys = sorted(set(da) | set(db), key=lambda k: (len(str(k)), str(k)))
    changed = ["%s %s→%s" % (k, fmt_dim(da.get(k)), fmt_dim(db.get(k))) for k in keys if da.get(k) != db.get(k)]
    if not changed:
        return []
    return ["- %s: %s%s" % (label, ", ".join(changed[:10]), " …" if len(changed) > 10 else "")]


def cell_range(coords):
    from openpyxl.utils import get_column_letter
    rows = [r for r, _ in coords]
    cols = [c for _, c in coords]
    return "%s%d:%s%d" % (get_column_letter(min(cols)), min(rows), get_column_letter(max(cols)), max(rows))


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="編集前後の Excel ブックの差分を出す")
    ap.add_argument("before")
    ap.add_argument("after")
    ap.add_argument("--limit", type=int, default=100)
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("openpyxl が必要: pip install openpyxl", file=sys.stderr)
        return 2
    for p in (args.before, args.after):
        if not Path(p).is_file():
            print("入力が存在しない: %s" % p, file=sys.stderr)
            return 2
    a = openpyxl.load_workbook(args.before)
    b = openpyxl.load_workbook(args.after)

    lines = []
    if a.sheetnames != b.sheetnames:
        lines.append("- シート: %s → %s" % (a.sheetnames, b.sheetnames))
    na = {k: v.attr_text for k, v in a.defined_names.items()}
    nb = {k: v.attr_text for k, v in b.defined_names.items()}
    if na != nb:
        changed = sorted(k for k in set(na) | set(nb) if na.get(k) != nb.get(k))
        lines.append("- 名前付き範囲: " + ", ".join("%s %s→%s" % (k, na.get(k), nb.get(k)) for k in changed))
    if lines:
        lines.append("")
    for name in a.sheetnames:
        if name not in b.sheetnames:
            continue
        wa, wb_ = a[name], b[name]
        head = []
        for label, va, vb in (
                ("表示状態", wa.sheet_state, wb_.sheet_state),
                ("使用範囲", wa.dimensions, wb_.dimensions),
                ("ウィンドウ枠の固定", wa.freeze_panes, wb_.freeze_panes),
                ("印刷範囲", wa.print_area, wb_.print_area),
                ("入力規則の数", len(wa.data_validations.dataValidation), len(wb_.data_validations.dataValidation)),
                ("条件付き書式の数", len(wa.conditional_formatting), len(wb_.conditional_formatting))):
            if va != vb:
                head.append("- %s: %s → %s" % (label, va, vb))
        ma = {str(r) for r in wa.merged_cells.ranges}
        mb = {str(r) for r in wb_.merged_cells.ranges}
        if ma != mb:
            head.append("- 結合セル: 消えた %s / 増えた %s" % (sorted(ma - mb) or "なし", sorted(mb - ma) or "なし"))
        head += diff_dims("列幅・非表示列", col_dims(wa), col_dims(wb_))
        head += diff_dims("行の高さ・非表示行", row_dims(wa), row_dims(wb_))

        changed = []
        # 実在するセルの和集合だけを比べる。片方にしか無いセルは cell() で既定のセルとして作って比べる
        for r, c in sorted(set(wa._cells) | set(wb_._cells)):
            ca, cb = wa.cell(r, c), wb_.cell(r, c)
            kinds = []
            if value_key(ca.value) != value_key(cb.value):
                kinds.append("値")
            ka, kb = style_key(ca), style_key(cb)
            if ka != kb:
                kinds.append("書式(%s)" % "・".join(style_names(ka, kb)))
            if kinds:
                changed.append(((r, c), ca.coordinate, "・".join(kinds), ca.value, cb.value))
        if head or changed:
            lines.append("## %s" % name)
            lines += head
            if changed:
                lines.append("- 変わったセル: %d 件（範囲 %s）" % (len(changed), cell_range([x[0] for x in changed])))
                lines.append("")
                lines.append("| セル | 種類 | 前 | 後 |")
                lines.append("|---|---|---|---|")
                for _, coord, kind, va, vb in changed[: args.limit]:
                    lines.append("| %s | %s | %s | %s |" % (coord, kind, show(va), show(vb)))
                if len(changed) > args.limit:
                    lines.append("")
                    lines.append("（先頭 %d 件のみ）" % args.limit)
            lines.append("")
    if not lines:
        print("差分なし")
        return 0
    print("\n".join(lines).rstrip())
    return 1


if __name__ == "__main__":
    sys.exit(main())
