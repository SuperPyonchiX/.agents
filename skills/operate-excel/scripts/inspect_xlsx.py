#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel ブックの構造を、中身を全件読まずに把握する。

    python inspect_xlsx.py <xlsx> [--sheet NAME] [--rows 10] [--cols 15]

出力（Markdown）:
    - ブック全体: シート一覧（非表示を含む）、名前付き範囲、openpyxl で保存すると失われる恐れのある要素
    - シートごと: 値のある範囲（書式だけのセルで使用範囲が広がっていればその旨）、数式セル数、
      結合セル、ウィンドウ枠の固定、印刷範囲、テーブル、入力規則・条件付き書式の数（0・なしの項目は省く）、
      先頭 --rows 行 × --cols 列のプレビュー（数式セルは「=式 → 保存されている計算結果」で表示）
    - シートが6枚以上あり --sheet を指定しないときは、プレビューを省いて各シートの概要だけを出す

終了コード:
    0  出力した
    1  開けない（壊れている、パスワード付き）
    2  引数の指定ミス、入力が存在しない、.xls など openpyxl が読めない形式、または openpyxl が無い

依存: openpyxl（pip install openpyxl）
"""

import argparse
import importlib.util
import re
import sys
import zipfile
from pathlib import Path

# zip 内のパターン → openpyxl で読み込んで保存したときに起きること
AT_RISK = [
    (r"^xl/vbaProject\.bin$", "マクロ（VBA）。keep_vba=True で開き .xlsm で保存しないと消える"),
    (r"^xl/pivotTables/", "ピボットテーブル。保持はされるが更新・編集はできない"),
    (r"^xl/slicers/|^xl/slicerCaches/", "スライサー。保存すると消える"),
    (r"^xl/timelines/", "タイムライン。保存すると消える"),
    (r"^xl/ctrlProps/", "フォームコントロール（ボタン・チェックボックス）。保存すると消える"),
    (r"^xl/externalLinks/", "外部ブックへのリンク。リンク先の値は更新されない"),
    (r"^xl/drawings/vmlDrawing", "コメントの書式・図形（VML）。書式が失われることがある"),
]
SHAPE_RE = re.compile(rb"<xdr:sp[ >]")
# シートの extLst は openpyxl が読まずに捨てる（読み込み時に「extension is not supported」と警告する）
SHEET_EXT = [
    (b"x14:dataValidation", "入力規則の拡張（別シートを参照するリストなど）"),
    (b"x14:conditionalFormatting", "条件付き書式の拡張（データバー・アイコンセットの新形式など）"),
    (b"x14:sparklineGroup", "スパークライン"),
]
PREVIEW_SHEETS = 5


def is_formula(v):
    if isinstance(v, str):
        return v.startswith("=")
    return type(v).__name__ in ("ArrayFormula", "DataTableFormula")


def formula_text(v):
    return v if isinstance(v, str) else (getattr(v, "text", None) or "=（データテーブル）")


def at_risk_items(path):
    items = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for pattern, message in AT_RISK:
            if any(re.search(pattern, n) for n in names):
                items.append(message)
        if any(n.startswith("xl/threadedComments/") for n in names):
            items.append("スレッド形式のコメント（返信つき）。保存すると返信のスレッドが消える")
        if any(n.startswith("xl/charts/chart") for n in names):
            items.append("グラフ。保持はされるが書式の一部が失われることがある")
        if any(n.startswith("xl/media/") for n in names) and importlib.util.find_spec("PIL") is None:
            items.append("画像。Pillow が入っていないと保存すると消える（pip install Pillow）")
        for n in names:
            if n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
                if SHAPE_RE.search(z.read(n)):
                    items.append("図形・テキストボックス（%s）。保存すると消える" % n)
            elif n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
                data = z.read(n)
                if b"extLst" in data:
                    found = [label for key, label in SHEET_EXT if key in data] or ["その他の拡張"]
                    items.append("%s（%s）。保存すると消える" % ("・".join(found), n))
    return items


def clip(s, limit):
    s = str(s).replace("\n", " ").replace("|", "\\|")
    return s if len(s) <= limit else s[: limit - 1] + "…"


def cell_text(v, cached, limit=40):
    if v is None:
        return ""
    if is_formula(v):
        # 長い式は式の側を切り、計算結果は残す
        result = clip("（保存値なし）" if cached is None else cached, limit // 2)
        return "%s → %s" % (clip(formula_text(v), max(limit - len(result) - 3, 10)), result)
    return clip(v, limit)


def data_bounds(ws):
    """値のあるセルの (最終行, 最終列)。ws.max_row は書式だけのセルでも広がるので使わない。"""
    max_r = max_c = 0
    for (r, c), cell in ws._cells.items():  # 実在するセルだけを走査する（使用範囲の全マスを作らない）
        if cell.value is not None:
            max_r, max_c = max(max_r, r), max(max_c, c)
    return max_r, max_c


def describe_sheet(ws, ws_values, rows, cols, preview):
    from openpyxl.utils import get_column_letter
    out = []
    state = "" if ws.sheet_state == "visible" else "（%s）" % ws.sheet_state
    out.append("## %s%s" % (ws.title, state))
    out.append("")
    data_r, data_c = data_bounds(ws)
    if data_r == 0:
        out.append("（空のシート）")
        return out
    used = "A1:%s%d（%d 行 × %d 列）" % (get_column_letter(data_c), data_r, data_r, data_c)
    if (ws.max_row, ws.max_column) != (data_r, data_c):
        used += "。書式だけのセルで使用範囲は %s まで広がっている" % ws.dimensions
    formulas = sum(1 for c in ws._cells.values() if is_formula(c.value))
    merged = [str(r) for r in ws.merged_cells.ranges]
    facts = [
        ("値のある範囲", used),
        ("数式セル", formulas),
        ("結合セル", merged and "%d 件 %s" % (len(merged), ", ".join(merged[:10]) + (" …" if len(merged) > 10 else ""))),
        ("ウィンドウ枠の固定", ws.freeze_panes),
        ("印刷範囲", ws.print_area),
        ("テーブル", ", ".join("%s(%s)" % (t.name, t.ref) for t in ws.tables.values())),
        ("入力規則", len(ws.data_validations.dataValidation)),
        ("条件付き書式", len(ws.conditional_formatting)),
    ]
    for k, v in facts:
        if v or k == "数式セル":
            out.append("- %s: %s" % (k, v))
    if not preview or rows <= 0:
        return out
    out.append("")

    max_r = min(data_r, rows)
    max_c = min(data_c, cols)
    header = ["行"] + [get_column_letter(c) for c in range(1, max_c + 1)]
    out.append("| " + " | ".join(header) + " |")
    out.append("|" + "---|" * len(header))
    # 計算結果は read_only で先頭の行だけ読む（シート全体を読み込まない）
    cached_rows = list(ws_values.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=True))
    for r in range(1, max_r + 1):
        cached = tuple(cached_rows[r - 1]) if r <= len(cached_rows) else ()
        cached += (None,) * (max_c - len(cached))
        cells = [cell_text(ws.cell(r, c).value, cached[c - 1]) for c in range(1, max_c + 1)]
        out.append("| %d | %s |" % (r, " | ".join(cells)))
    if data_r > rows or data_c > cols:
        out.append("")
        out.append("（先頭 %d 行 × %d 列のみ）" % (max_r, max_c))
    return out


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Excel ブックの構造を出す")
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", help="このシートだけ出す")
    ap.add_argument("--rows", type=int, default=10)
    ap.add_argument("--cols", type=int, default=15)
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("openpyxl が必要: pip install openpyxl", file=sys.stderr)
        return 2
    path = Path(args.xlsx)
    if not path.is_file():
        print("入力が存在しない: %s" % path, file=sys.stderr)
        return 2
    if path.suffix.lower() not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        print("openpyxl は %s を読めない。Excel か LibreOffice で .xlsx に変換してから扱う" % path.suffix,
              file=sys.stderr)
        return 2
    try:
        wb = openpyxl.load_workbook(path)
        wb_values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        risks = at_risk_items(path)
    except (zipfile.BadZipFile, OSError, KeyError) as e:
        print("開けない（壊れているかパスワード付き）: %s (%s)" % (path, e), file=sys.stderr)
        return 1

    lines = ["# %s" % path.name, ""]
    lines.append("- シート: " + ", ".join(
        ws.title + ("" if ws.sheet_state == "visible" else "（%s）" % ws.sheet_state) for ws in wb.worksheets))
    names = list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else []
    lines.append("- 名前付き範囲: " + (", ".join(names) if names else "なし"))
    if risks:
        lines.append("- **openpyxl で保存すると失われる・壊れる恐れのある要素**:")
        lines += ["  - " + r for r in risks]
    else:
        lines.append("- openpyxl で保存して失われる要素: 検出なし")
    lines.append("")

    if args.sheet and args.sheet not in wb.sheetnames:
        print("シートが無い: %s（あるのは %s）" % (args.sheet, ", ".join(wb.sheetnames)), file=sys.stderr)
        return 2
    targets = [wb[args.sheet]] if args.sheet else wb.worksheets
    preview = bool(args.sheet) or len(targets) <= PREVIEW_SHEETS
    for ws in targets:
        lines += describe_sheet(ws, wb_values[ws.title], args.rows, args.cols, preview)
        lines.append("")
    if not preview:
        lines.append("（シートが %d 枚あるのでプレビューを省いた。--sheet でシートを指定すると出る）" % len(targets))
    wb_values.close()
    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
