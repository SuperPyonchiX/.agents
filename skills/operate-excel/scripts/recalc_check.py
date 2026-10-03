#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel ブックの数式を再計算し、エラーになっているセルを洗い出す。

    python recalc_check.py <xlsx> [--engine auto|soffice|office] [--write-back]

openpyxl は数式の値を計算しないので、書いた数式が正しく評価されるかは表計算ソフトに計算させないと分からない。
このスクリプトはブックを一時フォルダにコピーし、コピーを再計算してから値を読む。元のファイルは変えない。
--write-back を付けると、Excel（office エンジン）で元のファイルを再計算して保存する
（計算結果を読む別のプログラムに渡すとき用。LibreOffice では書式が変わるので使えない）。

変換エンジン:
    soffice  LibreOffice で読み込み → xlsx で書き出すときに再計算させる
    office   Windows の MS Excel を PowerShell 経由の COM で操作し、CalculateFull してから保存する
    auto     soffice が見つかればそれ、無ければ Windows なら office（既定）

出力（JSON）:
    engine         使ったエンジン
    formulas       数式セルの数
    errors         エラーのセル（先頭50件）[{sheet, cell, error, formula}]
    error_summary  エラーの種類ごとの件数
    uncalculated   再計算後も値が空の数式セルの数。0 でなければ再計算が効いていない（engine を変えて試す）

終了コード:
    0  エラーなし（uncalculated も 0）
    1  エラーのセルがある、または uncalculated が 0 でない
    2  引数の指定ミス、入力が存在しない、または openpyxl が無い
    3  再計算に使えるエンジンが無い、または再計算に失敗した

依存: openpyxl（pip install openpyxl）。再計算には LibreOffice か Windows の MS Excel が要る
"""

import argparse
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!",
                "#SPILL!", "#CALC!", "#GETTING_DATA", "#FIELD!", "#BLOCKED!", "#UNKNOWN!"}
MAX_LISTED = 50

PS_EXCEL = r"""
$ErrorActionPreference = 'Stop'
$app = New-Object -ComObject Excel.Application
try {
  $app.AutomationSecurity = 3  # マクロを実行しない
  $app.Visible = $false
  $app.DisplayAlerts = $false
  $book = $app.Workbooks.Open($env:RECALC_PATH, 0, $false)
  try {
    $app.CalculateFull()
    $book.Save()
  } finally { $book.Close($false) }
} finally { $app.Quit() }
"""


def find_soffice():
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.isfile(c):
            return c
    return None


def recalc_office(path):
    env = dict(os.environ, RECALC_PATH=str(path))
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", PS_EXCEL],
                       capture_output=True, text=True, env=env, timeout=600)
    if r.returncode != 0:
        sys.stderr.write(r.stdout + r.stderr)
        return False
    return True


def recalc_soffice(soffice, path, workdir):
    """path を LibreOffice で読み直して書き出したファイルのパスを返す。"""
    outdir = Path(workdir) / "lo-out"
    outdir.mkdir()
    profile = Path(workdir) / "lo-profile"
    cmd = [soffice, "-env:UserInstallation=" + profile.as_uri(), "--headless",
           "--convert-to", "xlsx", "--outdir", str(outdir), str(path)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    out = outdir / (path.stem + ".xlsx")
    if r.returncode != 0 or not out.is_file():
        sys.stderr.write(r.stdout + r.stderr)
        return None
    return out


def scan(original, calculated):
    import openpyxl
    wb_f = openpyxl.load_workbook(original)
    wb_v = openpyxl.load_workbook(calculated, data_only=True)
    formulas = 0
    uncalculated = 0
    errors = []
    summary = Counter()
    for ws in wb_f.worksheets:
        if ws.title not in wb_v.sheetnames:
            continue
        wv = wb_v[ws.title]
        for row in ws.iter_rows():
            for c in row:
                is_formula = isinstance(c.value, str) and c.value.startswith("=")
                v = wv[c.coordinate].value
                if is_formula:
                    formulas += 1
                    if v is None:
                        uncalculated += 1
                if isinstance(v, str) and v in ERROR_VALUES:
                    summary[v] += 1
                    if len(errors) < MAX_LISTED:
                        errors.append({"sheet": ws.title, "cell": c.coordinate, "error": v,
                                       "formula": c.value if is_formula else None})
    return formulas, uncalculated, errors, summary


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="数式を再計算してエラーのセルを洗い出す")
    ap.add_argument("xlsx")
    ap.add_argument("--engine", choices=["auto", "soffice", "office"], default="auto")
    ap.add_argument("--write-back", action="store_true", help="Excel で元のファイルを再計算して保存する")
    args = ap.parse_args()

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("openpyxl が必要: pip install openpyxl", file=sys.stderr)
        return 2
    src = Path(args.xlsx).resolve()
    if not src.is_file():
        print("入力が存在しない: %s" % src, file=sys.stderr)
        return 2
    if src.suffix.lower() not in (".xlsx", ".xlsm"):
        print(".xlsx / .xlsm のみ対応", file=sys.stderr)
        return 2

    is_windows = platform.system() == "Windows" and shutil.which("powershell")
    soffice = find_soffice() if args.engine in ("auto", "soffice") and not args.write_back else None
    if args.write_back and not is_windows:
        print("--write-back は Windows の MS Excel でのみ使える", file=sys.stderr)
        return 3

    with tempfile.TemporaryDirectory() as work:
        if args.write_back:
            engine = "office"
            ok = recalc_office(src)
            calculated = src if ok else None
        elif soffice:
            engine = "soffice"
            calculated = recalc_soffice(soffice, src, work)
        elif args.engine in ("auto", "office") and is_windows:
            engine = "office"
            calculated = Path(work) / src.name
            shutil.copy2(src, calculated)
            if not recalc_office(calculated):
                calculated = None
        else:
            print("再計算に使えるエンジンが無い（LibreOffice も Windows の MS Excel も使えない）", file=sys.stderr)
            return 3
        if calculated is None:
            print("再計算に失敗した（engine=%s）" % engine, file=sys.stderr)
            return 3
        formulas, uncalculated, errors, summary = scan(src, calculated)

    result = {"engine": engine, "formulas": formulas, "errors": errors,
              "error_summary": dict(summary), "uncalculated": uncalculated}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if sum(summary.values()) > MAX_LISTED:
        print("エラーは %d 件。先頭 %d 件だけ列挙した" % (sum(summary.values()), MAX_LISTED), file=sys.stderr)
    return 1 if summary or uncalculated else 0


if __name__ == "__main__":
    sys.exit(main())
