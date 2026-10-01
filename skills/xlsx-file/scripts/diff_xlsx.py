#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""編集前後の Excel ブックを突き合わせ、変わったところを列挙する。

    python diff_xlsx.py <before.xlsx> <after.xlsx> [--limit 100]

既存ブックを編集したあと、意図したセル以外を変えていないかを確かめるために使う。
比べるもの: シートの集合と順序、各シートの使用範囲、結合セル、ウィンドウ枠の固定、
セルの値（数式は式の文字列で比べる）、セルの表示形式・フォント名・太字・塗りつぶし色。

出力（Markdown）: 構造の差分と、変わったセルの一覧（シートごと先頭 --limit 件）。

終了コード:
    0  差分なし
    1  差分あり（エラーではない。一覧が意図どおりかを確かめる）
    2  引数の指定ミス、入力が存在しない、または openpyxl が無い

依存: openpyxl（pip install openpyxl）
"""

import argparse
import sys
from pathlib import Path


def style_key(c):
    fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
    return (c.number_format, c.font.name if c.font else None, bool(c.font and c.font.b), fill)


def show(v):
    if v is None:
        return "（空）"
    s = str(v).replace("\n", " ").replace("|", "\\|")
    return s if len(s) <= 40 else s[:39] + "…"


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="編集前後の Excel ブックの差分を出す")
    ap.add_argument("before")
    ap.add_argument("after")
    ap.add_argument("--limit", type=int, default=100)
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("openpyxl が必要: pip install openpyxl", file=sys.stderr)
        return 2
    for p in (args.before, args.after):
        if not Path(p).is_file():
            print("入力が存在しない: %s" % p, file=sys.stderr)
            return 2
    a = openpyxl.load_workbook(args.before)
    b = openpyxl.load_workbook(args.after)

    lines = []
    if a.sheetnames != b.sheetnames:
        lines.append("- シート: %s → %s" % (a.sheetnames, b.sheetnames))
    for name in a.sheetnames:
        if name not in b.sheetnames:
            continue
        wa, wb_ = a[name], b[name]
        head = []
        if wa.dimensions != wb_.dimensions:
            head.append("- 使用範囲: %s → %s" % (wa.dimensions, wb_.dimensions))
        ma = {str(r) for r in wa.merged_cells.ranges}
        mb = {str(r) for r in wb_.merged_cells.ranges}
        if ma != mb:
            head.append("- 結合セル: 消えた %s / 増えた %s" % (sorted(ma - mb) or "なし", sorted(mb - ma) or "なし"))
        if wa.freeze_panes != wb_.freeze_panes:
            head.append("- ウィンドウ枠の固定: %s → %s" % (wa.freeze_panes, wb_.freeze_panes))
        rows = max(wa.max_row, wb_.max_row)
        cols = max(wa.max_column, wb_.max_column)
        changed = []
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                ca, cb = wa.cell(r, c), wb_.cell(r, c)
                kinds = []
                if ca.value != cb.value:
                    kinds.append("値")
                if style_key(ca) != style_key(cb):
                    kinds.append("書式")
                if kinds:
                    changed.append((ca.coordinate, "・".join(kinds), ca.value, cb.value))
        if head or changed:
            lines.append("## %s" % name)
            lines += head
            if changed:
                lines.append("- 変わったセル: %d 件" % len(changed))
                lines.append("")
                lines.append("| セル | 種類 | 前 | 後 |")
                lines.append("|---|---|---|---|")
                for coord, kind, va, vb in changed[: args.limit]:
                    lines.append("| %s | %s | %s | %s |" % (coord, kind, show(va), show(vb)))
                if len(changed) > args.limit:
                    lines.append("")
                    lines.append("（先頭 %d 件のみ）" % args.limit)
            lines.append("")
    if not lines:
        print("差分なし")
        return 0
    print("\n".join(lines))
    return 1


if __name__ == "__main__":
    sys.exit(main())
